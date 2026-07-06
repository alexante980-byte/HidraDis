import os
import sys
import numpy as np
import openpyxl
from datetime import datetime
import customtkinter as ctk
import tkinter as tk
from tkinter import filedialog, messagebox
from PIL import Image
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.drawing.image import Image as OXLImage
from openpyxl.utils import get_column_letter

# ------------------------------------------------------------------------------
# 1. PALETA DE COLORES Y CONSTANTES VISUALES
# ------------------------------------------------------------------------------
COLOR_FONDO_MODULOS = "#D6E4F0"
COLOR_CONTENEDOR    = "#FFFFFF"
COLOR_BORDE_TECNICO = "#B1D0E0"
COLOR_ACCENTO       = "#1A374D"
COLOR_ACCENTO_SUAVE = "#406882"
COLOR_RESULTADO     = "#d32f2f"
COLOR_TEAL          = "#5BC0BE"
COLOR_TEAL_TITULOS  = "#0E7C86"
FUENTE_BASE         = "Segoe UI"

ctk.set_appearance_mode("light")

# ------------------------------------------------------------------------------
# 2. EXCEPCIONES PERSONALIZADAS DE DISEÑO NORMADO
# ------------------------------------------------------------------------------
class SAFValidationError(Exception):
    """Excepción para detener el cálculo si no se cumplen los límites normativos SAF:
      · 1.7 ≤ Fr₁ ≤ 17       — rango de resalto estable (FEMA, 2010)
      · V₁ ≤ 18.3 m/s (60 ft/s) — límite de cavitación en bloques de rápida (FEMA, 2010)"""
    pass

class USBRVIValidationError(Exception):
    """Excepción para detener el cálculo si no se cumplen los límites normativos
    del cuenco de impacto USBR Tipo VI:
      Q ≤ 11.33 m³/s  |  V ≤ 15.24 m/s  |  1.1 ≤ Fr ≤ 10.0  (FEMA, 2010)."""
    pass
# ------------------------------------------------------------------------------
# 3. MOTOR HIDRÁULICO COMÚN
# ------------------------------------------------------------------------------
def resolver_ruta(ruta_relativa: str) -> str:
    try:
        base_path = sys._MEIPASS
    except AttributeError:
        base_path = os.path.abspath(".")
    return os.path.join(base_path, ruta_relativa)

def calcular_presion_vapor(temperatura_c: float) -> float:
    p_mmhg = 10 ** (8.07131 - (1730.63 / (temperatura_c + 233.426)))
    return 0.0136 * p_mmhg

def calcular_presion_atmosferica(altitud_msnm: float) -> float:
    return 10.33 - (altitud_msnm / 900.0)

def cargar_imagen_en(parent, path, size, side, refs):
    if os.path.exists(path):
        img = ctk.CTkImage(light_image=Image.open(path), size=size)
        refs.append(img)
        ctk.CTkLabel(parent, image=img, text="").pack(side=side, padx=5)

# ------------------------------------------------------------------------------
# 4. UTILIDADES VISUALES
# ------------------------------------------------------------------------------
def centrar_ventana(ventana: ctk.CTk, ancho: int, alto: int) -> None:
    ventana.update_idletasks()
    pantalla_w = ventana.winfo_screenwidth()
    pantalla_h = ventana.winfo_screenheight()
    pos_x = max(0, (pantalla_w - ancho) // 2)
    pos_y = max(0, (pantalla_h - alto) // 2)
    ventana.geometry(f"{ancho}x{alto}+{pos_x}+{pos_y}")

# ------------------------------------------------------------------------------
# 5. VISTA: PANTALLA DE BIENVENIDA
# ------------------------------------------------------------------------------
class PantallaInicio(ctk.CTkFrame):
    def __init__(self, master, control_navegacion, ancho_ref=1450, alto_ref=920, **kwargs):
        super().__init__(master, **kwargs)
        self.configure(fg_color="#0F172A")
        self._img_refs = []

        hdr = ctk.CTkFrame(self, fg_color="#1E293B", corner_radius=0)
        hdr.pack(fill="x")

        hl = ctk.CTkFrame(hdr, fg_color="transparent")
        hl.pack(fill="x", padx=35, pady=12)

        cargar_imagen_en(hl, resolver_ruta(os.path.join("Imagenes", "Escuela_Politécnica_Nacional.png")),
                          (68, 68), "left", self._img_refs)
        cargar_imagen_en(hl, resolver_ruta(os.path.join("Imagenes", "logo_fica.png")),
                          (135, 43), "right", self._img_refs)

        tc = ctk.CTkFrame(hl, fg_color="transparent")
        tc.pack(side="left", expand=True)
        ctk.CTkLabel(tc, text="ESCUELA POLITÉCNICA NACIONAL",
                     font=(FUENTE_BASE, 16, "bold"), text_color="#F8FAFC").pack()
        ctk.CTkLabel(tc, text="FACULTAD DE INGENIERÍA CIVIL Y AMBIENTAL",
                     font=(FUENTE_BASE, 12, "bold"), text_color="#38BDF8").pack(pady=1)
        ctk.CTkLabel(tc, text="TRABAJO DE INTEGRACIÓN CURRICULAR",
                     font=(FUENTE_BASE, 10, "italic"), text_color="#94A3B8").pack()

        ctk.CTkFrame(self, fg_color="#0EA5E9", height=4, corner_radius=0).pack(fill="x")

        title_area = ctk.CTkFrame(self, fg_color="transparent")
        title_area.pack(pady=(40, 0))

        ctk.CTkLabel(title_area, text="HIDRA DIS",
                     font=(FUENTE_BASE, 46, "bold"), text_color="#F8FAFC").pack()
        ctk.CTkLabel(title_area,
                     text="Software Integrado para el Dimensionamiento Geométrico y Evaluación Hidráulica de Cuencos Disipadores de Energía",
                     font=(FUENTE_BASE, 14, "italic"), text_color="#38BDF8").pack(pady=4)

        ctk.CTkFrame(self, fg_color="#334155", height=1, corner_radius=0).pack(
            fill="x", padx=100, pady=(15, 20))

        ctk.CTkLabel(self,
                     text=(
                         "Generación sistemática de memorias de cálculo, dimensionamiento geométrico y evaluación\n"
                         "de riesgo por cavitación en obras de descarga bajo normativas internacionales\n"
                         "del cuenco SAF y USBR tipo VI."
                     ),
                     font=(FUENTE_BASE, 13), text_color="#E2E8F0",
                     justify="center").pack(pady=(0, 35))

        cards_frame = ctk.CTkFrame(self, fg_color="transparent")
        cards_frame.pack(pady=10)

        saf_card = ctk.CTkFrame(cards_frame, fg_color="#1E293B", corner_radius=16,
                                 border_width=1, border_color="#38BDF8",
                                 width=350, height=210)
        saf_card.grid(row=0, column=0, padx=30)
        saf_card.pack_propagate(False)

        badge_saf = ctk.CTkFrame(saf_card, fg_color="#0EA5E9", corner_radius=8,
                                  width=120, height=32)
        badge_saf.pack(pady=(20, 6))
        badge_saf.pack_propagate(False)
        ctk.CTkLabel(badge_saf, text="CUENCO SAF",
                     font=(FUENTE_BASE, 12, "bold"), text_color="#0F172A").pack(expand=True)

        ctk.CTkLabel(saf_card, text="Saint Anthony Falls Laboratory",
                     font=(FUENTE_BASE, 11, "italic"), text_color="#7DD3FC").pack(pady=(0, 4))
        ctk.CTkLabel(saf_card, text="Azud y Cuenco",
                     font=(FUENTE_BASE, 11), text_color="#94A3B8").pack(pady=(0, 15))

        ctk.CTkButton(saf_card, text="Ir al módulo SAF  →",
                      font=(FUENTE_BASE, 12, "bold"), width=220, height=36,
                      fg_color="#0284C7", hover_color="#0369A1", corner_radius=8,
                      text_color="#FFFFFF",
                      command=control_navegacion.mostrar_saf).pack(pady=4)

        vi_card = ctk.CTkFrame(cards_frame, fg_color="#1E293B", corner_radius=16,
                                border_width=1, border_color="#38BDF8",
                                width=350, height=210)
        vi_card.grid(row=0, column=1, padx=30)
        vi_card.pack_propagate(False)

        badge_vi = ctk.CTkFrame(vi_card, fg_color="#0EA5E9", corner_radius=8,
                                 width=140, height=32)
        badge_vi.pack(pady=(20, 6))
        badge_vi.pack_propagate(False)
        ctk.CTkLabel(badge_vi, text="CUENCO USBR VI",
                     font=(FUENTE_BASE, 12, "bold"), text_color="#0F172A").pack(expand=True)

        ctk.CTkLabel(vi_card, text="Bureau of Reclamation — Tipo VI",
                     font=(FUENTE_BASE, 11, "italic"), text_color="#7DD3FC").pack(pady=(0, 4))
        ctk.CTkLabel(vi_card, text="Sección rectangular y circular",
                     font=(FUENTE_BASE, 10), text_color="#94A3B8").pack(pady=(0, 15))

        ctk.CTkButton(vi_card, text="Ir al módulo USBR VI  →",
                      font=(FUENTE_BASE, 12, "bold"), width=220, height=36,
                      fg_color="#0284C7", hover_color="#0369A1", corner_radius=8,
                      text_color="#FFFFFF",
                      command=control_navegacion.mostrar_usbr_vi).pack(pady=4)

        footer_ini = ctk.CTkFrame(self, fg_color="#0F172A", corner_radius=0, height=40)
        footer_ini.pack(fill="x", side="bottom")
        footer_ini.pack_propagate(False)
        ctk.CTkLabel(footer_ini,
                     text=f"Autor: Alex Ante   |   Ingeniería Civil |   {datetime.now().strftime('%d/%m/%Y')}",
                     font=(FUENTE_BASE, 11), text_color="#64748B").pack(expand=True)

# ------------------------------------------------------------------------------
# 6. MÓDULO DEL CUENCO TIPO SAF
# ------------------------------------------------------------------------------
class InterfaceSAF(ctk.CTkFrame):
    """Módulo de dimensionamiento del cuenco disipador tipo SAF."""

    def __init__(self, master, control_navegacion, **kwargs):
        super().__init__(master, **kwargs)
        self.configure(fg_color=COLOR_FONDO_MODULOS)
        self._img_refs = []
        self.inputs = {}
        self.input_frames = {}
        self.results = {}
        self.result_name_labels = {}
        self.result_unit_labels   = {}
        self.status_labels = {}
        self.data_export = {}
        self.control_navegacion = control_navegacion

        self.ruta_epn        = resolver_ruta(os.path.join("Imagenes", "Escuela_Politécnica_Nacional.png"))
        self.ruta_fica       = resolver_ruta(os.path.join("Imagenes", "logo_fica.png"))
        self.ruta_planta     = resolver_ruta(os.path.join("Imagenes", "Vista en planta SAF.png"))
        self.ruta_elevacion  = resolver_ruta(os.path.join("Imagenes", "Vista en elevación_SAF.png"))
        self.ruta_isometrica = resolver_ruta(os.path.join("Imagenes", "Vista isométrica_SAF.png"))

        self.init_module_ui()

    def init_module_ui(self):
        header = ctk.CTkFrame(self, fg_color=COLOR_FONDO_MODULOS, corner_radius=0)
        header.pack(fill="x", padx=0, pady=0)
        content_head = ctk.CTkFrame(header, fg_color="transparent")
        content_head.pack(pady=5, fill="x")

        self.btn_volver = ctk.CTkButton(
            content_head, text="⬅ Menú Principal", width=130, height=30,
            fg_color=COLOR_ACCENTO, hover_color=COLOR_ACCENTO_SUAVE, font=(FUENTE_BASE, 11, "bold"),
            command=self.control_navegacion.mostrar_inicio)
        self.btn_volver.pack(side="left", padx=15)

        self.load_img(content_head, self.ruta_epn, (45, 45), "left")
        titles_frame = ctk.CTkFrame(content_head, fg_color="transparent")
        titles_frame.pack(side="left", padx=15, expand=True)
        ctk.CTkLabel(titles_frame, text="DISEÑO DE CUENCO DISIPADOR TIPO SAF",
                     font=(FUENTE_BASE, 17, "bold"), text_color=COLOR_ACCENTO).pack()
        ctk.CTkLabel(titles_frame,
                     text="Cálculo y dimensionamiento geométrico de cuencos disipadores SAF para un óptimo desempeño hidráulico en obras de derivación.",
                     font=(FUENTE_BASE, 10, "italic"), text_color=COLOR_ACCENTO_SUAVE).pack()
        self.load_img(content_head, self.ruta_fica, (110, 35), "right")

        body = ctk.CTkFrame(self, fg_color=COLOR_FONDO_MODULOS)
        body.pack(fill="both", expand=True, padx=10, pady=2)

        type_frame = ctk.CTkFrame(body, fg_color=COLOR_FONDO_MODULOS, corner_radius=8,
                                   border_width=1, border_color=COLOR_BORDE_TECNICO)
        type_frame.pack(fill="x", pady=(0, 4))
        ctk.CTkLabel(type_frame, text="Tipo de Proyecto:",
                     font=(FUENTE_BASE, 11, "bold")).pack(side="left", padx=10, pady=3)
        self.tipo_var = ctk.StringVar(value="Cuenco")
        ctk.CTkRadioButton(type_frame, text="Azud (Cuenco ancho constante)",
                           variable=self.tipo_var, value="Derivacion",
                           command=self.toggle_inputs,
                           font=(FUENTE_BASE, 10)).pack(side="left", padx=10)
        ctk.CTkRadioButton(type_frame, text="Cuenco (ancho variable)",
                           variable=self.tipo_var, value="Cuenco",
                           command=self.toggle_inputs,
                           font=(FUENTE_BASE, 10)).pack(side="left", padx=10)

        panels = ctk.CTkFrame(body, fg_color="transparent")
        panels.pack(fill="both", expand=True)

        left_side = ctk.CTkFrame(panels, fg_color=COLOR_FONDO_MODULOS, width=680)
        left_side.pack(side="left", fill="both", expand=False, padx=(0, 5))
        left_side.pack_propagate(False)
        self.left_scroll = ctk.CTkScrollableFrame(left_side, fg_color=COLOR_FONDO_MODULOS, label_text="")
        self.left_scroll.pack(fill="both", expand=True)

        right_side = ctk.CTkFrame(panels, fg_color=COLOR_FONDO_MODULOS, width=710)
        right_side.pack(side="right", fill="both", expand=True)
        right_side.pack_propagate(False)
        self.right_scroll = ctk.CTkScrollableFrame(right_side, fg_color=COLOR_FONDO_MODULOS)
        self.right_scroll.pack(fill="both", expand=True)

        self.setup_sections()
        self.setup_graphics()

        footer = ctk.CTkFrame(self, fg_color=COLOR_FONDO_MODULOS, height=42, corner_radius=0,
                               border_width=1, border_color=COLOR_BORDE_TECNICO)
        footer.pack(fill="x", side="bottom")
        footer.pack_propagate(False)

        info_tic = ctk.CTkFrame(footer, fg_color="transparent")
        info_tic.pack(side="left", padx=20)
        ctk.CTkLabel(info_tic, text="Autor: Alex Ante", height=16,
                     font=(FUENTE_BASE, 9, "bold"), text_color=COLOR_ACCENTO).pack(anchor="w", pady=0)
        ctk.CTkLabel(info_tic, text=datetime.now().strftime("%d/%m/%Y"), height=14,
                     font=(FUENTE_BASE, 8), text_color=COLOR_ACCENTO_SUAVE).pack(anchor="w", pady=0)

        btn_frame = ctk.CTkFrame(footer, fg_color="transparent")
        btn_frame.pack(side="right", padx=20, pady=6)
        ctk.CTkButton(btn_frame, text="Calcular", command=self.calcular,
                      fg_color=COLOR_ACCENTO, width=90, height=30,
                      font=(FUENTE_BASE, 11, "bold")).pack(side="left", padx=4)
        ctk.CTkButton(btn_frame, text="Limpiar", command=self.limpiar,
                      fg_color="#698396", width=80, height=30,
                      font=(FUENTE_BASE, 11)).pack(side="left", padx=4)
        ctk.CTkButton(btn_frame, text="Exportar Excel", command=self.exportar_excel,
                      fg_color="#2E7D32", width=120, height=30,
                      font=(FUENTE_BASE, 11, "bold")).pack(side="left", padx=4)

    def load_img(self, parent, path, size, side):
        cargar_imagen_en(parent, path, size, side, self._img_refs)

    def toggle_inputs(self):
        if self.tipo_var.get() == "Derivacion":
            # Ocultar campos de Cuenco
            self.input_frames["D0_frame"].grid_forget()
            self.input_frames["V1_frame"].grid_forget()
            self.input_frames["Fr1_frame"].grid_forget()
            self.input_frames["WB_frame"].grid_forget()
            self.input_frames["y1c_frame"].grid_forget()
            # Mostrar campos de Azud 
            self.input_frames["Br_frame"].grid(row=0, column=1, padx=6, pady=4)
            self.input_frames["y1_frame"].grid(row=1, column=0, padx=6, pady=4)
            self.input_frames["Temp_frame"].grid(row=1, column=1, padx=6, pady=4)
            self.input_frames["Alt_frame"].grid(row=2,  column=0, padx=6, pady=4)
            if "LA" in self.results:
                self.results["LA"].master.grid_remove()
            # Restaurar etiquetas de Fr₁ y V₁ en resultados
            if "Fr1_ap" in self.result_name_labels:
                self.result_name_labels["Fr1_ap"].configure(text="Número Froude [Fr₁]:")
                self.result_unit_labels["Fr1_ap"].configure(text="-")
            if "V1" in self.result_name_labels:
                self.result_name_labels["V1"].configure(text="Velocidad Flujo [V₁]:")
                self.result_unit_labels["V1"].configure(text="m/s")
        else:  # Cuenco (ancho variable) 
            # Ocultar campo exclusivo de Azud
            self.input_frames["Br_frame"].grid_forget()
            self.input_frames["y1c_frame"].grid_forget()
            self.input_frames["WB_frame"].grid_forget()
            # Mostrar campos de cuenco 
            self.input_frames["D0_frame"].grid( row=0, column=1, padx=6, pady=4)
            self.input_frames["y1_frame"].grid( row=1, column=0, padx=6, pady=4)
            self.input_frames["V1_frame"].grid( row=1, column=1, padx=6, pady=4)
            self.input_frames["Fr1_frame"].grid(row=2, column=0, padx=6, pady=4)
            self.input_frames["Temp_frame"].grid(row=2, column=1, padx=6, pady=4)
            self.input_frames["Alt_frame"].grid( row=3, column=0, padx=6, pady=4)
            if "LA" in self.results:
                self.results["LA"].master.grid()
            # Fr₁ y V₁ son datos de entrada en este modo 
            if "Fr1_ap" in self.result_name_labels:
                self.result_name_labels["Fr1_ap"].configure(text="")
                self.result_unit_labels["Fr1_ap"].configure(text="")
                self.results["Fr1_ap"].configure(text="")
            if "V1" in self.result_name_labels:
                self.result_name_labels["V1"].configure(text="")
                self.result_unit_labels["V1"].configure(text="")
                self.results["V1"].configure(text="")

    def setup_sections(self):
        ctk.CTkLabel(self.left_scroll, text="Datos:",
                     font=(FUENTE_BASE, 13, "bold"), text_color=COLOR_ACCENTO,
                     anchor="w").pack(fill="x", padx=5, pady=(2, 0))
        in_f = ctk.CTkFrame(self.left_scroll, fg_color=COLOR_FONDO_MODULOS, corner_radius=6,
                             border_width=1, border_color=COLOR_BORDE_TECNICO)
        in_f.pack(fill="x", pady=(1, 5))

        grid = ctk.CTkFrame(in_f, fg_color="transparent")
        grid.pack(padx=5, pady=4, anchor="w")

        self.input_frames["Q_frame"]    = self.create_input(grid, "Caudal (Q):",                    "m³/s", "Q",    row=0, col=0)
        self.input_frames["Br_frame"]   = self.create_input(grid, "Ancho azud (Bᵣ):",             "m",    "Br",   row=0, col=1)
        self.input_frames["D0_frame"]   = self.create_input(grid, "Ancho de ingreso (Bᵣ):",        "m",    "D0",   row=0, col=1)
        self.input_frames["y1_frame"]   = self.create_input(grid, "Tirante supercrítico (y₁):",   "m",    "y1",   row=1, col=0)
        self.input_frames["y1c_frame"]  = self.create_input(grid, "Prof. contraída entrada (y₁):", "m",   "y1c",  row=1, col=0)
        self.input_frames["V1_frame"]   = self.create_input(grid, "Velocidad ingreso (V₁):",      "m/s",  "V1_in",row=1, col=0)
        self.input_frames["Fr1_frame"]  = self.create_input(grid, "Froude ingreso (Fr₁):",        "-",    "Fr1_in",row=2, col=1)
        self.input_frames["WB_frame"]   = self.create_input(grid, "Ancho del cuenco (WB):",       "m",    "WB_in",row=3, col=0)
        self.input_frames["Temp_frame"] = self.create_input(grid, "Temperatura (T):",              "°C",   "Temp", row=1, col=1)
        self.input_frames["Alt_frame"]  = self.create_input(grid, "Altitud (A):",                  "msnm", "Alt",  row=2, col=0)

        self.toggle_inputs()

        ctk.CTkLabel(self.left_scroll, text="Resultados:",
                     font=(FUENTE_BASE, 13, "bold"), text_color=COLOR_ACCENTO,
                     anchor="w").pack(fill="x", padx=5, pady=(2, 0))
        res_f = ctk.CTkFrame(self.left_scroll, fg_color=COLOR_FONDO_MODULOS, corner_radius=6,
                              border_width=1, border_color=COLOR_BORDE_TECNICO)
        res_f.pack(fill="both", expand=True, pady=(1, 5))

        sections = [
            ("1. Parámetros de Cuenco", [
                ("Ancho Cuenco", "WB", "WB", "m"),
                ("Ancho Rápida", "WB₁", "WB1", "m"),
                ("Ancho Intermedio", "WB₂", "WB2", "m")]),
            ("2. Parámetros Hidráulicos", [
                ("Número Froude",     "Fr₁", "Fr1_ap", "-"),
                ("Velocidad Flujo",   "V₁",  "V1",     "m/s"),
                ("Tirante Crítico",   "yc",  "yc_r",   "m"),
                ("Tirante Conjugado", "y₂",  "d2",     "m")]),
            ("3. Condiciones del cuenco SAF", [
                ("Factor Corrección", "C", "C", "-")]),
            ("4. Bloques de Rápida", [
                ("Número Bloques", "n꜀♭", "ncb", "u"),
                ("Ancho", "w꜀♭", "wcb", "m"),
                ("Espaciamiento", "a꜀♭", "acb", "m"),
                ("Altura", "h꜀♭", "hcb", "m")]),
            ("5. Bloques de Fondo", [
                ("Número Bloques", "nբ♭", "nfb", "u"),
                ("Ancho", "wբ♭", "wfb", "m"),
                ("Espaciamiento", "aբ♭", "afb", "m"),
                ("Altura", "hբ♭", "hfb", "m"),
                ("Espesor Cresta", "e", "e", "m"),
                ("Distancia Pared", "dᵣₑₐₗ", "dreal", "m")]),
            ("6. Aproximación Geométrica", [
                ("Transición Ent.", "Lₐ", "LA", "m"),
                ("Longitud Cuenco", "LB", "LB", "m"),
                ("Ubicación Bloques", "Dₑₛ", "Des", "m")]),
            ("7. Estructura Final", [
                ("Altura Umbral", "h₄", "h4", "m"),
                ("Altura Pared", "h₆", "h6", "m"),
                ("Borde Libre", "z", "z_bl", "m"),
                ("Ancho Salida", "WB₃", "WB3", "m")]),
            ("8. Cavitación (Ingreso de la rápida)", [
                ("Índice de cavitación", "σ", "sigma", "-")])
        ]

        for title, items in sections:
            self.create_result_group(res_f, title, items)

    def create_input(self, parent, label, unit, key, row, col):
        f = ctk.CTkFrame(parent, fg_color="transparent")
        f.grid(row=row, column=col, padx=8, pady=4, sticky="w")
        ctk.CTkLabel(f, text=label, width=130, font=(FUENTE_BASE, 10), anchor="w").pack(side="left")
        e = ctk.CTkEntry(f, width=65, height=22, fg_color=COLOR_CONTENEDOR)
        e.pack(side="left", padx=1)
        e.bind("<Return>", lambda _event: self.calcular())
        ctk.CTkLabel(f, text=unit, font=(FUENTE_BASE, 9), text_color="gray",
                     width=32, anchor="w").pack(side="left", padx=2)
        self.inputs[key] = e
        return f

    def create_result_group(self, parent, title, items):
        ctk.CTkLabel(parent, text=title, font=(FUENTE_BASE, 11, "bold"),
                     text_color=COLOR_TEAL_TITULOS, anchor="w").pack(fill="x", padx=10, pady=(5, 1))
        container = ctk.CTkFrame(parent, fg_color="transparent")
        container.pack(fill="x", padx=8, pady=1)

        for col_idx in range(3):
            container.grid_columnconfigure(col_idx, weight=1, minsize=215)

        for i, item in enumerate(items):
            r, c = divmod(i, 3)
            cell = ctk.CTkFrame(container, fg_color="transparent")
            cell.grid(row=r, column=c, padx=3, pady=2, sticky="w")

            _name_lbl = ctk.CTkLabel(cell, text=f"{item[0]} [{item[1]}]:",
                         width=100, anchor="w", font=(FUENTE_BASE, 9.5))
            _name_lbl.pack(side="left")
            self.result_name_labels[item[2]] = _name_lbl
            res_box = ctk.CTkLabel(cell, text="0.00", width=62, height=19,
                                    fg_color=COLOR_CONTENEDOR, corner_radius=3,
                                    font=(FUENTE_BASE, 9.5, "bold"), text_color=COLOR_RESULTADO)
            res_box.pack(side="left", padx=1)
            self.results[item[2]] = res_box
            _unit_lbl = ctk.CTkLabel(cell, text=item[3], width=32, text_color="gray",
                              font=(FUENTE_BASE, 9), anchor="w")
            _unit_lbl.pack(side="left")
            self.result_unit_labels[item[2]] = _unit_lbl

        if "SAF" in title:
            # Cuadro 1: verificación Froude
            sf = ctk.CTkFrame(container, fg_color=COLOR_CONTENEDOR, corner_radius=6,
                               border_width=1, border_color=COLOR_BORDE_TECNICO)
            sf.grid(row=0, column=1, padx=3, pady=2, sticky="w")
            ctk.CTkLabel(sf, text="Froude  1.7 < Fr₁ < 17", font=(FUENTE_BASE, 9),
                         text_color=COLOR_ACCENTO_SUAVE).pack(pady=(3, 0), padx=10)
            self.status_labels["Fr1"] = ctk.CTkLabel(
                sf, text="--", font=(FUENTE_BASE, 10, "bold"), text_color="gray")
            self.status_labels["Fr1"].pack(pady=(0, 3), padx=10)
            # Cuadro 2: verificación velocidad SAF 
            sf_v = ctk.CTkFrame(container, fg_color=COLOR_CONTENEDOR, corner_radius=6,
                                 border_width=1, border_color=COLOR_BORDE_TECNICO)
            sf_v.grid(row=0, column=2, padx=3, pady=2, sticky="w")
            ctk.CTkLabel(sf_v, text="Velocidad  V₁ ≤ 18.3 m/s", font=(FUENTE_BASE, 9),
                         text_color=COLOR_ACCENTO_SUAVE).pack(pady=(3, 0), padx=10)
            self.status_labels["V1_lim"] = ctk.CTkLabel(
                sf_v, text="--", font=(FUENTE_BASE, 10, "bold"), text_color="gray")
            self.status_labels["V1_lim"].pack(pady=(0, 3), padx=10)

        elif "Fondo" in title:
            sf = ctk.CTkFrame(container, fg_color="transparent")
            sf.grid(row=2, column=0, columnspan=3, sticky="w", pady=1)
            ctk.CTkLabel(sf, text="Ocupación (40-55% Bᵣ):",
                         font=(FUENTE_BASE, 9)).pack(side="left", padx=(2, 0))
            self.results["Oc"] = ctk.CTkLabel(
                sf, text="0.0%", width=48, font=(FUENTE_BASE, 9, "bold"), text_color=COLOR_RESULTADO)
            self.results["Oc"].pack(side="left", padx=3)
            self.status_labels["Oc"] = ctk.CTkLabel(
                sf, text="--", font=(FUENTE_BASE, 9, "bold"), text_color="gray")
            self.status_labels["Oc"].pack(side="left", padx=3)

        elif "Cavitación" in title:
            sf = ctk.CTkFrame(container, fg_color="transparent")
            sf.grid(row=1, column=0, columnspan=3, sticky="w", pady=1)
            ctk.CTkLabel(sf, text="Diagnóstico Cavitación:",
                         font=(FUENTE_BASE, 9, "bold")).pack(side="left", padx=(2, 0))
            self.status_labels["Cav"] = ctk.CTkLabel(
                sf, text="Evaluar diseño", font=(FUENTE_BASE, 9, "bold"), text_color="gray")
            self.status_labels["Cav"].pack(side="left", padx=5)

    def setup_graphics(self):
        ctk.CTkLabel(self.right_scroll, text="Esquema:",
                     font=(FUENTE_BASE, 13, "bold"), text_color=COLOR_ACCENTO,
                     anchor="w").pack(fill="x", padx=5, pady=(2, 2))

        formatos = {
            "Vista en Planta":    (500, 300),
            "Vista en Elevación": (500, 300),
            "Vista Isométrica":   (500, 300)
        }

        for title, path in [
            ("Vista en Planta",    self.ruta_planta),
            ("Vista en Elevación", self.ruta_elevacion),
            ("Vista Isométrica",   self.ruta_isometrica)
        ]:
            f = ctk.CTkFrame(self.right_scroll, fg_color=COLOR_FONDO_MODULOS, corner_radius=6,
                              border_width=1, border_color=COLOR_BORDE_TECNICO)
            f.pack(fill="x", pady=3, padx=2)
            ctk.CTkLabel(f, text=title, font=(FUENTE_BASE, 10, "bold"),
                         text_color=COLOR_ACCENTO).pack(pady=1)

            if os.path.exists(path):
                w, h = formatos[title]
                img = ctk.CTkImage(light_image=Image.open(path), size=(w, h))
                self._img_refs.append(img)
                ctk.CTkLabel(f, image=img, text="").pack(pady=2, padx=10)
            else:
                ctk.CTkLabel(f, text=f"[{title} ",
                             font=(FUENTE_BASE, 9, "italic"), text_color="gray").pack(pady=20)

            ctk.CTkLabel(f, text="Adaptado Chow (1994)",
                         font=(FUENTE_BASE, 8, "italic"),
                         text_color="#555555").pack(pady=(0, 2), anchor="e", padx=10)

    def calcular(self):
        try:
            g = 9.81
            mode = self.tipo_var.get()

            Q        = float(self.inputs["Q"].get().replace(',', '.'))
            temp_agua= float(self.inputs["Temp"].get().replace(',', '.'))
            altitud  = float(self.inputs["Alt"].get().replace(',', '.'))

            if mode == "Derivacion":
                Br  = float(self.inputs["Br"].get().replace(',', '.'))
                y1  = float(self.inputs["y1"].get().replace(',', '.'))
                V1  = Q / (Br * y1)
                Fr1 = V1 / np.sqrt(g * y1)
                Wb1 = Br
                Wb  = Br
                LA  = 0
            else:  # Cuenco (ancho variable) — todos los datos son entrada directa
                Br  = float(self.inputs["D0"].get().replace(',', '.'))     # ancho de ingreso [m]
                y1  = float(self.inputs["y1"].get().replace(',', '.'))     # calado supercrítico [m]
                V1  = float(self.inputs["V1_in"].get().replace(',', '.'))  # velocidad [m/s]
                Fr1 = float(self.inputs["Fr1_in"].get().replace(',', '.')) # Froude [-]
                # Ancho sección de rápida
                Wb1 = Br if V1 < 6.0 else 2.5 * Br
                # Ancho del cuenco
                WB_min = 1.7 * Q / (np.sqrt(g) * Br**1.5)
                Wb  = max(Wb1, WB_min)
                # Longitud de transición de entrada — expansión lateral 1:1 (45°)
                LA  = max(0.0, (Wb - Br) / 2.0)

            # Fr₁ y V₁ fuera de rango: se calculan igual y se advierten al final

            yc  = ((Q / Wb)**2 / g)**(1/3)
            d2  = (y1 / 2) * (np.sqrt(1 + 8 * Fr1**2) - 1)

            if   1.7 < Fr1 <= 5.5:  C_val = 1.1 - (Fr1**2 / 120)
            elif 5.5 < Fr1 <= 11:   C_val = 0.85
            else:                    C_val = 1.0 - (Fr1**2 / 800)

            LB  = (4.5 * d2) / (C_val * (Fr1**0.76))
            Wb2 = Wb

            ncb  = max(1, int(Wb1 / (1.5 * y1))) if (1.5 * y1) > 0 else 1
            wcb  = Wb1 / (2 * ncb) if ncb > 0 else 0
            nfb  = ncb
            wfb  = Wb2 / (2 * nfb) if nfb > 0 else 0

            Hatm      = calcular_presion_atmosferica(altitud)
            Hv        = calcular_presion_vapor(temp_agua)
            carga_vel = (V1**2) / (2 * g)
            sigma     = (Hatm + y1 - Hv) / carga_vel if carga_vel > 0 else 0

            if   sigma > 1.0:        estado_cav, color_cav = "NO HAY CAVITACIÓN  (Condición segura)", "green"
            elif 0.4 <= sigma <= 1.0: estado_cav, color_cav = "RIESGO MODERADO  (Revisar diseño)", "orange"
            else:                    estado_cav, color_cav = "ALTO RIESGO  (Rediseñar estructura)", "red"

            # y1_r y yc_r se muestran siempre en sección 2 con etiquetas fijas

            res_map = {
                "WB": Wb, "WB1": Wb1, "WB2": Wb2, "V1": V1,
                "Fr1_ap": Fr1, "Fr1": Fr1, "d2": d2, "C": C_val,
                "LA": LA, "LB": LB, "Des": LB / 3,
                "ncb": ncb, "wcb": wcb, "acb": wcb, "hcb": y1,
                "nfb": nfb, "wfb": wfb, "afb": wfb, "hfb": y1,
                "e": 0.5 * y1,
                "dreal": (Wb2 - (nfb * wfb) - (nfb - 1) * wfb) / 2 if nfb > 0 else 0,
                "h4":  (0.07 * (d2 / C_val)),
                "h6": d2 * (1 + (1 / (3 * C_val))),
                "z_bl": d2 / 3,
                "WB3": Br if mode == "Derivacion" else Wb1 + 2 * LB / Br,
                "sigma": sigma,
                "yc_r": yc,   # sección 2 — tirante crítico
            }

            for k, v in res_map.items():
                if k in self.results:
                    self.results[k].configure(
                        text=f"{int(v)}" if k in ["ncb", "nfb"]
                        else f"{v:.2f}" if k in ["sigma", "Fr1_ap", "yc_r", "V1", "C"]
                        else f"{v:.2f}" if k in ["d2"]
                        else f"{v:.2f}")
            if mode == "Cuenco":
                # Fr₁ y V₁ son datos de entrada 
                self.result_name_labels["Fr1_ap"].configure(text="")
                self.result_unit_labels["Fr1_ap"].configure(text="")
                self.results["Fr1_ap"].configure(text="")
                self.result_name_labels["V1"].configure(text="")
                self.result_unit_labels["V1"].configure(text="")
                self.results["V1"].configure(text="")
            else:  # Azud — Fr₁ y V₁ 
                self.result_name_labels["Fr1_ap"].configure(text="Número Froude [Fr₁]:")
                self.result_unit_labels["Fr1_ap"].configure(text="-")
                self.result_name_labels["V1"].configure(text="Velocidad Flujo [V₁]:")
                self.result_unit_labels["V1"].configure(text="m/s")

            ocup = (nfb * wfb / Wb2) * 100 if nfb > 0 else 0
            self.results["Oc"].configure(text=f"{ocup:.1f}%")
            self.status_labels["Oc"].configure(
                text="CUMPLE" if 40 <= ocup <= 55 else "REVISAR",
                text_color="green" if 40 <= ocup <= 55 else "red")
            self.status_labels["Cav"].configure(text=estado_cav, text_color=color_cav)
            # Verificaciones normativas — advertencia sin bloquear el cálculo
            avisos = []

            if Fr1 < 1.7 or Fr1 > 17.0:
                self.status_labels["Fr1"].configure(text="EXCEDE", text_color="red")
                avisos.append(
                    f"• Número de Froude (Fr₁ = {Fr1:.2f}) fuera del rango SAF 1.7 – 17 (FEMA, 2010).\n"
                    f"  Las ecuaciones de dimensionamiento SAF no son válidas fuera de este rango."
                )
            else:
                self.status_labels["Fr1"].configure(text="CUMPLE", text_color="green")

            if V1 > 18.3:
                self.status_labels["V1_lim"].configure(text="EXCEDE", text_color="red")
                avisos.append(
                    f"• Velocidad de ingreso (V₁ = {V1:.2f} m/s) supera 18.3 m/s — 60 ft/s (FEMA, 2010).\n"
                    f"  Por encima de este umbral puede producirse cavitación severa en los bloques."
                )
            else:
                self.status_labels["V1_lim"].configure(text="CUMPLE", text_color="green")

            if avisos:
                messagebox.showwarning(
                    "Advertencia — Límites Normativos SAF",
                    "Los resultados se han calculado y se muestran, pero se detectaron las "
                    "siguientes condiciones fuera de norma (FEMA, 2010):\n\n"
                    + "\n\n".join(avisos)
                    + "\n\nSe recomienda revisar el diseño antes de usar la memoria de cálculo."
                )

            self.data_export = {
                "inputs": {
                    "Tipo Proyecto": mode, "Q": Q, "y1": y1,
                    "Temp": temp_agua, "Alt": altitud,
                    "Br": Br,
                    "V1_calc": V1, "Fr1_calc": Fr1},
                "results": res_map, "cav_text": estado_cav
            }

        except SAFValidationError as val_err:
            self.limpiar()
            messagebox.showwarning("Límite de Diseño Excedido", str(val_err))
        except Exception as err:
            messagebox.showerror("Error de Cálculo",
                                 f"Verifique las variables numéricas.\nDetalle: {err}")

    def limpiar(self):
        for e in self.inputs.values():
            e.delete(0, tk.END)
        for r in self.results.values(): r.configure(text="0.00")
        for s in self.status_labels.values(): s.configure(text="--", text_color="gray")
        if "Oc" in self.results: self.results["Oc"].configure(text="0.0%")
        if "Cav" in self.status_labels:
            self.status_labels["Cav"].configure(text="Evaluar diseño", text_color="gray")
        self.data_export = {}
        # Restaurar etiquetas y unidades según modo activo
        self.toggle_inputs()

    def exportar_excel(self):
        if not self.data_export:
            messagebox.showwarning("Atención", "Por favor, realice el cálculo antes de exportar.")
            return
        path = filedialog.asksaveasfilename(defaultextension=".xlsx",
                                            filetypes=[("Archivos Excel", "*.xlsx")])
        if not path: return
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Memoria SAF"
            ws.views.sheetView[0].showGridLines = True

            fill_header  = PatternFill(start_color="1A374D", end_color="1A374D", fill_type="solid")
            fill_section = PatternFill(start_color="406882", end_color="406882", fill_type="solid")
            fill_zebra   = PatternFill(start_color="F9FBFC", end_color="F9FBFC", fill_type="solid")

            font_title   = Font(name=FUENTE_BASE, size=15, bold=True, color="1A374D")
            font_section = Font(name=FUENTE_BASE, size=11, bold=True, color="FFFFFF")
            font_header  = Font(name=FUENTE_BASE, size=10, bold=True, color="FFFFFF")
            font_bold    = Font(name=FUENTE_BASE, size=10, bold=True)
            font_regular = Font(name=FUENTE_BASE, size=10)

            thin = Border(
                left=Side(style='thin', color='B1D0E0'), right=Side(style='thin', color='B1D0E0'),
                top=Side(style='thin', color='B1D0E0'), bottom=Side(style='thin', color='B1D0E0'))

            ws["A1"] = "MEMORIA TÉCNICA DE DISEÑO HIDRÁULICO - CUENCO SAF"
            ws["A1"].font = font_title
            ws["A4"] = "1. DATOS DE ENTRADA"
            ws.merge_cells("A4:D4")
            ws["A4"].fill = fill_section
            ws["A4"].font = font_section

            headers = ["Parámetro", "Variable", "Valor", "Unidad"]
            for c_idx, h in enumerate(headers, start=1):
                cell = ws.cell(row=5, column=c_idx, value=h)
                cell.fill = fill_header; cell.font = font_header
                cell.alignment = Alignment(horizontal="center")

            inp = self.data_export["inputs"]
            datos_in = [
                ("Tipo de Estructura de Entrada", "Proyecto", inp["Tipo Proyecto"], "-"),
                ("Caudal del Diseño",             "Q",        inp["Q"],             "m³/s"),
                ("Tirante Supercrítico",          "y₁",       inp["y1"],            "m"),
                ("Temperatura del Fluido",        "T",        inp["Temp"],          "°C"),
                ("Altitud del Proyecto",          "A",        inp["Alt"],           "msnm"),
                ("Ancho de Ingreso",              "Bᵣ",       inp["Br"],            "m"),
                ("Velocidad",                     "V₁",       inp["V1_calc"],       "m/s"),
                ("Número de Froude",              "Fr₁",      inp["Fr1_calc"],      "-"),
            ]

            r_curr = 6
            for item in datos_in:
                ws.cell(row=r_curr, column=1, value=item[0]).font = font_regular
                ws.cell(row=r_curr, column=2, value=item[1]).font = font_bold
                v_cell = ws.cell(row=r_curr, column=3, value=item[2]); v_cell.font = font_regular
                if isinstance(item[2], float): v_cell.number_format = "0.00"
                ws.cell(row=r_curr, column=4, value=item[3]).font = font_regular
                for col in range(1, 5): ws.cell(row=r_curr, column=col).border = thin
                r_curr += 1

            r_curr += 1
            ws.cell(row=r_curr, column=1, value="2. RESULTADOS DEL DIMENSIONAMIENTO").font = font_section
            ws.merge_cells(start_row=r_curr, start_column=1, end_row=r_curr, end_column=4)
            ws.cell(row=r_curr, column=1).fill = fill_section

            r_curr += 1
            for c_idx, h in enumerate(headers, start=1):
                cell = ws.cell(row=r_curr, column=c_idx, value=h)
                cell.fill = fill_header; cell.font = font_header
                cell.alignment = Alignment(horizontal="center")

            r_curr += 1
            res = self.data_export["results"]
            datos_out = [
                ("Ancho del Cuenco",            "WB",   res["WB"],    "m"),
                ("Ancho Rápida Entrada",        "WB₁",  res["WB1"],   "m"),
                ("Ancho Sección Central",       "WB₂",  res["WB2"],   "m"),
                ("Ancho Salida",                "WB₃",  res["WB3"],   "m"),
                ("Velocidad",                   "V₁",   res["V1"],    "m/s"),
                ("Número de Froude",            "Fr₁",  res["Fr1"],   "-"),
                ("Tirante Crítico",             "y꜀",   res["yc_r"],  "m"),
                ("Tirante Conjugado",           "y₂",   res["d2"],    "m"),
                ("Factor Corrección",            "C",    res["C"],     "-"),
                *([("Transición de Entrada",    "Lₐ",   res["LA"],    "m")]
                  if self.data_export["inputs"]["Tipo Proyecto"] == "Cuenco" else []),
                ("Longitud Cuenco SAF",         "L_B",  res["LB"],    "m"),
                ("Ubicación Bloques",           "Dₑₛ",  res["Des"],   "m"),
                ("N.° Bloques Rápida",          "n꜀♭",  res["ncb"],   "u"),
                ("Ancho Bloques Rápida",        "w꜀♭",  res["wcb"],   "m"),
                ("Espaciamiento Rápida",        "a꜀♭",  res["acb"],   "m"),
                ("Altura Bloques Rápida",       "h꜀♭",  res["hcb"],   "m"),
                ("N.° Bloques Fondo",           "nբ♭",  res["nfb"],   "u"),
                ("Ancho Bloques Fondo",         "wբ♭",  res["wfb"],   "m"),
                ("Espaciamiento Fondo",         "aբ♭",  res["afb"],   "m"),
                ("Altura Bloques Fondo",        "hբ♭",  res["hfb"],   "m"),
                ("Espesor Cresta",              "e",    res["e"],     "m"),
                ("Distancia a Pared",           "dᵣₑₐₗ",res["dreal"],"m"),
                ("Altura Umbral de Salida",     "h₄",   res["h4"],    "m"),
                ("Altura de Paredes",           "h₆",   res["h6"],    "m"),
                ("Borde Libre",                 "z",    res["z_bl"],  "m"),
                ("Índice de Cavitación",        "σ",    res["sigma"], "-"),
                ("Diagnóstico Cavitación", "Estado", self.data_export["cav_text"], "—"),
            ]

            for item in datos_out:
                ws.cell(row=r_curr, column=1, value=item[0]).font = font_regular
                ws.cell(row=r_curr, column=2, value=item[1]).font = font_bold
                v_cell = ws.cell(row=r_curr, column=3, value=item[2])
                if isinstance(item[2], float):
                    v_cell.font = font_regular
                    v_cell.number_format = "0.00"
                elif isinstance(item[2], int):
                    v_cell.font = font_regular; v_cell.number_format = "0"
                else:
                    v_cell.font = font_bold
                ws.cell(row=r_curr, column=4, value=item[3]).font = font_regular
                if r_curr % 2 == 0:
                    for col in range(1, 5): ws.cell(row=r_curr, column=col).fill = fill_zebra
                for col in range(1, 5): ws.cell(row=r_curr, column=col).border = thin
                r_curr += 1

            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                ws.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 3, 11)
            ws.column_dimensions['A'].width = 36
            ws.column_dimensions['C'].width = 45

            # Esquemas técnicos SAF 
            r_curr += 2
            ws.cell(row=r_curr, column=1, value="3. ESQUEMAS TÉCNICOS DEL CUENCO SAF").font = font_section
            ws.merge_cells(start_row=r_curr, start_column=1, end_row=r_curr, end_column=4)
            ws.cell(row=r_curr, column=1).fill = fill_section
            r_curr += 2
            _esquemas_saf = [
                ("Vista en Planta SAF",    self.ruta_planta),
                ("Vista en Elevación SAF", self.ruta_elevacion),
                ("Vista Isométrica SAF",   self.ruta_isometrica),
            ]
            from PIL import Image as _PILImg
            for etiqueta, ruta_img in _esquemas_saf:
                if os.path.exists(ruta_img):
                    ws.cell(row=r_curr, column=1, value=etiqueta).font = font_bold
                    r_curr += 1
                    try:
                        _pi = _PILImg.open(ruta_img)
                        _ow, _oh = _pi.size
                        _tw = 550
                        _th = max(60, int(_oh * _tw / _ow))
                        _xi = OXLImage(ruta_img)
                        _xi.width  = _tw
                        _xi.height = _th
                        ws.add_image(_xi, f"A{r_curr}")
                        _rows = (_th // 18) + 2
                        for _rr in range(r_curr, r_curr + _rows):
                            ws.row_dimensions[_rr].height = 18
                        r_curr += _rows + 1
                    except Exception:
                        pass

            wb.save(path)
            messagebox.showinfo("Éxito", "Memoria técnica exportada correctamente a Excel.")
        except Exception as e:
            messagebox.showerror("Error", f"No se pudo guardar.\nDetalle: {e}")

# ------------------------------------------------------------------------------
# 7. MÓDULO DEL CUENCO TIPO USBR VI
# ------------------------------------------------------------------------------
class InterfaceUSBRVI(ctk.CTkFrame):
    """Módulo de dimensionamiento del cuenco de impacto USBR Tipo VI."""

    def __init__(self, master, control_navegacion, **kwargs):
        super().__init__(master, **kwargs)
        self.configure(fg_color=COLOR_FONDO_MODULOS)
        self.control_navegacion = control_navegacion
        self.inputs  = {}
        self.results = {}
        self.checks  = {}
        self.status_labels = {}
        self.datos_para_exportar = None
        self._img_refs = []

        self.ruta_epn      = resolver_ruta(os.path.join("Imagenes", "Escuela_Politécnica_Nacional.png"))
        self.ruta_fica     = resolver_ruta(os.path.join("Imagenes", "logo_fica.png"))
        self.ruta_planta   = resolver_ruta(os.path.join("Imagenes", "Vista en planta_USBR_VI.png"))
        self.ruta_elevacion= resolver_ruta(os.path.join("Imagenes", "Sección A-A_USBR_VI.png"))

        self.init_module_ui()

    def init_module_ui(self):
        header = ctk.CTkFrame(self, fg_color=COLOR_FONDO_MODULOS, corner_radius=0)
        header.pack(fill="x", padx=0, pady=0)
        content_head = ctk.CTkFrame(header, fg_color="transparent")
        content_head.pack(pady=5, fill="x")

        self.btn_volver = ctk.CTkButton(
            content_head, text="⬅ Menú Principal", width=130, height=30,
            fg_color=COLOR_ACCENTO, hover_color=COLOR_ACCENTO_SUAVE, font=(FUENTE_BASE, 11, "bold"),
            command=self.control_navegacion.mostrar_inicio)
        self.btn_volver.pack(side="left", padx=15)

        self.load_img(content_head, self.ruta_epn, (45, 45), "left")
        titles_frame = ctk.CTkFrame(content_head, fg_color="transparent")
        titles_frame.pack(side="left", padx=15, expand=True)
        ctk.CTkLabel(titles_frame, text="DISEÑO DE CUENCO DISIPADOR USBR TIPO VI",
                     font=(FUENTE_BASE, 17, "bold"), text_color=COLOR_ACCENTO).pack()
        ctk.CTkLabel(titles_frame,
                     text="Cálculo y dimensionamiento geométrico de cuencos disipadores por impacto USBR Tipo VI",
                     font=(FUENTE_BASE, 10, "italic"), text_color=COLOR_ACCENTO_SUAVE).pack()
        self.load_img(content_head, self.ruta_fica, (110, 35), "right")

        body = ctk.CTkFrame(self, fg_color=COLOR_FONDO_MODULOS)
        body.pack(fill="both", expand=True, padx=10, pady=2)

        self.tipo_sec = tk.IntVar(value=1)
        type_frame = ctk.CTkFrame(body, fg_color=COLOR_FONDO_MODULOS, corner_radius=8,
                                   border_width=1, border_color=COLOR_BORDE_TECNICO)
        type_frame.pack(fill="x", pady=(0, 4))
        ctk.CTkLabel(type_frame, text="Tipo de Sección:",
                     font=(FUENTE_BASE, 11, "bold")).pack(side="left", padx=10, pady=3)
        ctk.CTkRadioButton(type_frame, text="Rectangular", variable=self.tipo_sec, value=1,
                           command=self.toggle_entries,
                           font=(FUENTE_BASE, 10)).pack(side="left", padx=10)
        ctk.CTkRadioButton(type_frame, text="Circular", variable=self.tipo_sec, value=2,
                           command=self.toggle_entries,
                           font=(FUENTE_BASE, 10)).pack(side="left", padx=10)

        panels = ctk.CTkFrame(body, fg_color="transparent")
        panels.pack(fill="both", expand=True)
        panels.grid_columnconfigure(0, weight=3)
        panels.grid_columnconfigure(1, weight=2)
        panels.grid_rowconfigure(0, weight=1)

        left_side = ctk.CTkFrame(panels, fg_color=COLOR_FONDO_MODULOS)
        left_side.grid(row=0, column=0, sticky="nsew", padx=(0, 5))
        self.left_scroll = ctk.CTkScrollableFrame(left_side, fg_color=COLOR_FONDO_MODULOS)
        self.left_scroll.pack(fill="both", expand=True)

        right_side = ctk.CTkFrame(panels, fg_color=COLOR_FONDO_MODULOS)
        right_side.grid(row=0, column=1, sticky="nsew")
        self.right_scroll = ctk.CTkScrollableFrame(right_side, fg_color=COLOR_FONDO_MODULOS)
        self.right_scroll.pack(fill="both", expand=True)

        ctk.CTkLabel(self.left_scroll, text="Datos:",
                     font=(FUENTE_BASE, 13, "bold"),
                     text_color=COLOR_ACCENTO).pack(anchor="w", padx=2, pady=(2, 0))
        self.setup_inputs()

        ctk.CTkLabel(self.left_scroll, text="Resultados:",
                     font=(FUENTE_BASE, 13, "bold"),
                     text_color=COLOR_ACCENTO).pack(anchor="w", padx=2, pady=(4, 0))
        self.setup_results_table()

        ctk.CTkLabel(self.right_scroll, text="Esquema:",
                     font=(FUENTE_BASE, 13, "bold"),
                     text_color=COLOR_ACCENTO).pack(anchor="w", padx=5, pady=(2, 0))
        self.setup_graphics()

        footer = ctk.CTkFrame(self, fg_color=COLOR_FONDO_MODULOS, height=42, corner_radius=0,
                               border_width=1, border_color=COLOR_BORDE_TECNICO)
        footer.pack(fill="x", side="bottom")
        footer.pack_propagate(False)

        info_frame = ctk.CTkFrame(footer, fg_color="transparent")
        info_frame.pack(side="left", padx=25)
        ctk.CTkLabel(info_frame, text="Autor: Alex Ante", height=16,
                     font=(FUENTE_BASE, 9, "bold"), text_color=COLOR_ACCENTO).pack(anchor="w", pady=0)
        ctk.CTkLabel(info_frame, text=datetime.now().strftime("%d/%m/%Y"), height=14,
                     font=(FUENTE_BASE, 8), text_color=COLOR_ACCENTO_SUAVE).pack(anchor="w", pady=0)

        btn_frame = ctk.CTkFrame(footer, fg_color="transparent")
        btn_frame.pack(side="right", padx=25, pady=6)
        ctk.CTkButton(btn_frame, text="Calcular", command=self.calcular,
                      fg_color=COLOR_ACCENTO, width=95, height=30,
                      font=(FUENTE_BASE, 11, "bold")).pack(side="left", padx=4)
        ctk.CTkButton(btn_frame, text="Limpiar", command=self.limpiar,
                      fg_color="#698396", width=85, height=30).pack(side="left", padx=4)
        ctk.CTkButton(btn_frame, text="Exportar Excel", command=self.exportar,
                      fg_color="#2E7D32", width=125, height=30,
                      font=(FUENTE_BASE, 11, "bold")).pack(side="left", padx=4)

    def load_img(self, parent, path, size, side):
        cargar_imagen_en(parent, path, size, side, self._img_refs)

    def setup_inputs(self):
        frame = ctk.CTkFrame(self.left_scroll, fg_color=COLOR_FONDO_MODULOS, corner_radius=8,
                              border_width=1, border_color=COLOR_BORDE_TECNICO)
        frame.pack(fill="x", pady=2)

        grid_in = ctk.CTkFrame(frame, fg_color="transparent")
        grid_in.pack(fill="x", padx=10, pady=6)
        for _ci in range(3):
            grid_in.grid_columnconfigure(_ci, weight=1)

        fields_info = [
            ("Caudal",      "Q", "m³/s", 0, 0),
            ("n Manning",   "n", "-",    0, 1),
            ("Pendiente",   "S", "m/m",  0, 2),
            ("Ancho base",  "b", "m",    1, 0),
            ("Diámetro",    "D", "m",    1, 1),
            ("Temperatura", "T", "°C",   1, 2),
            ("Altitud",     "A", "msnm", 2, 0),
        ]

        for name, sym, unit, row, col in fields_info:
            cell = ctk.CTkFrame(grid_in, fg_color="transparent")
            cell.grid(row=row, column=col, padx=5, pady=2, sticky="w")
            ctk.CTkLabel(cell, text=f"{name} ({sym}):", width=95,
                         anchor="w", font=(FUENTE_BASE, 10.5)).pack(side="left")
            ent = ctk.CTkEntry(cell, width=65, height=22, fg_color="white", justify="center")
            ent.pack(side="left", padx=2)
            ent.bind("<Return>", lambda _event: self.calcular())
            self.inputs[sym] = ent
            ctk.CTkLabel(cell, text=unit, width=35, anchor="w",
                         font=(FUENTE_BASE, 9.5), text_color="gray").pack(side="left")

        self.toggle_entries()

    def toggle_entries(self):
        for k in ["Q", "n", "S", "T", "A"]: self.inputs[k].configure(state="normal")
        if self.tipo_sec.get() == 1:
            self.inputs["b"].configure(state="normal",   fg_color="white")
            self.inputs["D"].configure(state="disabled", fg_color="#E0E0E0")
        else:
            self.inputs["b"].configure(state="disabled", fg_color="#E0E0E0")
            self.inputs["D"].configure(state="normal",   fg_color="white")

    def setup_results_table(self):
        master_frame = ctk.CTkFrame(self.left_scroll, fg_color=COLOR_FONDO_MODULOS, corner_radius=8,
                                     border_width=1, border_color=COLOR_BORDE_TECNICO)
        master_frame.pack(fill="both", expand=True, pady=2, ipady=4)

        ctk.CTkLabel(master_frame, text="Verificaciones Normativas",
                     font=(FUENTE_BASE, 11, "bold"), text_color=COLOR_ACCENTO).pack(anchor="w", padx=12, pady=(4, 1))
        check_container = ctk.CTkFrame(master_frame, fg_color="transparent")
        check_container.pack(fill="x", padx=12, pady=(1, 4))
        for col_idx in range(3):
            check_container.grid_columnconfigure(col_idx, weight=1)

        for col_idx, (k, text) in enumerate([
            ("V",  "Velocidad  V ≤ 15.24 m/s"),
            ("Q",  "Caudal  Q ≤ 11.33 m³/s"),
            ("Fr", "Froude  1.1 ≤ Fr ≤ 10.0")
        ]):
            celda = ctk.CTkFrame(check_container, fg_color=COLOR_CONTENEDOR, corner_radius=6,
                                  border_width=1, border_color=COLOR_BORDE_TECNICO)
            celda.grid(row=0, column=col_idx, padx=4, pady=2, sticky="nsew")
            ctk.CTkLabel(celda, text=text, font=(FUENTE_BASE, 9.5),
                         text_color=COLOR_ACCENTO_SUAVE).pack(pady=(4, 0))
            self.checks[k] = ctk.CTkLabel(celda, text="---",
                                           font=(FUENTE_BASE, 10.5, "bold"), text_color="gray")
            self.checks[k].pack(pady=(0, 4))

        res_config = [
            ("1. Flujo de Aproximación", [
                ("Tirante normal", "Yn", "m"),
                ("Área mojada",    "A",  "m²"),
                ("Velocidad",      "V",  "m/s"),
                ("Número Froude",  "Fr", "-")
            ]),
            ("2. Geometría Cuenco USBR VI", [
                ("Ancho mínimo",  "Winf", "m"),
                ("Ancho máximo",  "Wsup", "m"),
                ("Ancho cuenco",    "W",  "m"),
                ("Longitud cuenco", "L",  "m"),
                ("Espesor frontal", "f",  "m"),
                ("Transición",      "e",  "m"),
                ("Altura total",    "H",  "m"),
                ("Lado bloque",     "a",  "m"),
                ("Ancho deflector", "b",  "m"),
                ("Contrafuerte",    "c",  "m"),
                ("Espesor pantalla","t",  "m"),
                ("Diámetro Roca",   "Dr", "m")
            ]),
            ("3. Cavitación (Ingreso de la rápida)", [
                ("Índice cavitación", "σ", "-")
            ])
        ]

        for sec_name, items in res_config:
            ctk.CTkLabel(master_frame, text=sec_name,
                         font=(FUENTE_BASE, 11, "bold"), text_color=COLOR_TEAL_TITULOS,
                         anchor="w").pack(fill="x", padx=12, pady=(6, 1))
            container = ctk.CTkFrame(master_frame, fg_color="transparent")
            container.pack(fill="x", padx=12, pady=1)

            for col_idx in range(3):
                container.grid_columnconfigure(col_idx, weight=1, minsize=250)

            for i, (name, sym, unit) in enumerate(items):
                r, c = divmod(i, 3)
                cell = ctk.CTkFrame(container, fg_color="transparent")
                cell.grid(row=r, column=c, padx=3, pady=2, sticky="w")

                ctk.CTkLabel(cell, text=f"{name} [{sym}]:", width=118,
                             anchor="w", font=(FUENTE_BASE, 10)).pack(side="left")
                res_box = ctk.CTkLabel(cell, text="0.00", width=68, height=20,
                                        fg_color="white", corner_radius=4,
                                        font=(FUENTE_BASE, 10.5, "bold"), text_color=COLOR_RESULTADO)
                res_box.pack(side="left", padx=2)
                self.results[sym] = res_box
                ctk.CTkLabel(cell, text=unit, width=42, anchor="w",
                             font=(FUENTE_BASE, 9.5), text_color="gray").pack(side="left")

            if "Cavitación" in sec_name:
                cav_f = ctk.CTkFrame(master_frame, fg_color="transparent")
                cav_f.pack(fill="x", padx=12, pady=(2, 4))
                ctk.CTkLabel(cav_f, text="Diagnóstico de Cavitación:",
                             font=(FUENTE_BASE, 10, "bold"), text_color=COLOR_ACCENTO).pack(side="left")
                self.status_labels["Cav"] = ctk.CTkLabel(
                    cav_f, text="Evaluar diseño numérico",
                    font=(FUENTE_BASE, 10, "bold"), text_color="gray")
                self.status_labels["Cav"].pack(side="left", padx=10)

    def setup_graphics(self):
        frame = ctk.CTkFrame(self.right_scroll, fg_color=COLOR_FONDO_MODULOS, corner_radius=8,
                              border_width=1, border_color=COLOR_BORDE_TECNICO)
        frame.pack(fill="both", expand=True, pady=2)

        for title, path in [
            ("VISTA EN PLANTA",               self.ruta_planta),
            ("VISTA EN ELEVACIÓN / SECCIÓN",  self.ruta_elevacion)
        ]:
            ctk.CTkLabel(frame, text=title, font=(FUENTE_BASE, 11, "bold"),
                         text_color=COLOR_ACCENTO).pack(pady=(6, 1))
            self.render_image(frame, path)
            ctk.CTkLabel(frame, text="Adaptado Beichley (1971)",
                         font=(FUENTE_BASE, 8, "italic"),
                         text_color="#555555").pack(pady=(0, 6), anchor="e", padx=15)

    def render_image(self, parent, path):
        try:
            if os.path.exists(path):
                ctk_img = ctk.CTkImage(light_image=Image.open(path), size=(460, 210))
                self._img_refs.append(ctk_img)
                ctk.CTkLabel(parent, text="", image=ctk_img).pack(pady=2, padx=10)
            else:
                ctk.CTkLabel(parent, text="[Esquema no encontrado en recursos]",
                             height=100, width=420, text_color="gray",
                             font=(FUENTE_BASE, 9, "italic"),
                             fg_color="#E0E0E0", corner_radius=6).pack(pady=5)
        except Exception:
            pass

    def get_geo_logic(self, y_val, tipo, b_reg, D_reg):
        if tipo == 1:
            A = b_reg * y_val; P = b_reg + 2 * y_val; T = b_reg; theta = None
        else:
            y_s   = max(1e-6, min(y_val, D_reg - 1e-6))
            theta = 2 * np.arccos(1 - (2 * y_s / D_reg))
            A     = (D_reg**2 / 8) * (theta - np.sin(theta))
            P     = (D_reg / 2) * theta
            T     = D_reg * np.sin(theta / 2)
        R = A / P if P > 0 else 0
        return A, P, R, T, theta

    def f_manning(self, y_val, Q_reg, n_reg, S_reg, tipo, b_reg, D_reg):
        A, _, R, _, _ = self.get_geo_logic(y_val, tipo, b_reg, D_reg)
        if A <= 0: return -Q_reg
        return (1 / n_reg) * A * (R**(2/3)) * np.sqrt(S_reg) - Q_reg

    def calcular(self):

        if self.tipo_sec.get() == 0:
            messagebox.showwarning("Atención", "Seleccione el tipo de sección.")
            return
        try:
            Q        = float(self.inputs["Q"].get().replace(',', '.'))
            n_m      = float(self.inputs["n"].get().replace(',', '.'))
            S_m      = float(self.inputs["S"].get().replace(',', '.'))
            temp_agua= float(self.inputs["T"].get().replace(',', '.'))
            altitud  = float(self.inputs["A"].get().replace(',', '.'))

            t_sec = self.tipo_sec.get()
            b_s   = float(self.inputs["b"].get().replace(',', '.')) if t_sec == 1 else 0
            D_d   = float(self.inputs["D"].get().replace(',', '.')) if t_sec == 2 else 0
            g     = 9.81

            y_calc = 0.2 * D_d if t_sec == 2 else 0.5; tol = 1e-9
            for _ in range(200):
                h_step = 1e-8
                f  = self.f_manning(y_calc, Q, n_m, S_m, t_sec, b_s, D_d)
                df = (self.f_manning(y_calc + h_step, Q, n_m, S_m, t_sec, b_s, D_d) - f) / h_step
                if abs(df) < 1e-12: df = 1e-12
                paso   = 0.5 * (f / df)
                y_new  = y_calc - paso
                if t_sec == 2: y_new = max(1e-4, min(y_new, 0.85 * D_d))
                else:          y_new = max(1e-3, y_new)
                if abs(y_new - y_calc) < tol: y_calc = y_new; break
                y_calc = y_new

            y_norm = round(y_calc, 2)
            A_f, P_f, R_f, T_f, _ = self.get_geo_logic(y_norm, t_sec, b_s, D_d)
            V_t = Q / A_f

            Fr = V_t / np.sqrt(g * (A_f / T_f))

            v_cumple  = V_t <= 15.24
            q_cumple  = Q   <= 11.33
            fr_cumple = 1.1 <= Fr <= 10.0

            self.checks["V"].configure(
                text="SÍ CUMPLE" if v_cumple else "NO CUMPLE",
                text_color="green" if v_cumple else "red")
            self.checks["Q"].configure(
                text="SÍ CUMPLE" if q_cumple else "NO CUMPLE",
                text_color="green" if q_cumple else "red")
            self.checks["Fr"].configure(
                text="SÍ CUMPLE" if fr_cumple else "NO CUMPLE",
                text_color="green" if fr_cumple else "red")

            # Límites fuera de norma: se continúa el cálculo y se advierte al final
            Q_cfs  = Q * 35.3147
            W_inf_ft = 1.47 * (Q_cfs**0.4)           # límite inferior — FEMA (2010)
            W_sup_ft = 1.79 * (Q_cfs**0.4)           # límite superior — FEMA (2010)
            W_ft   = (W_inf_ft + W_sup_ft) / 2       # ancho de diseño (valor promedio)
            m_conv = 0.3048

            H_atm    = calcular_presion_atmosferica(altitud)
            H_v      = calcular_presion_vapor(temp_agua)
            carga_vel= (V_t**2) / (2 * g)
            sigma_val= (H_atm + y_norm - H_v) / carga_vel if carga_vel > 0 else 0

            if   sigma_val >= 1.0: estado_cav, color_cav = "Condición Segura (No hay Cavitación)", "green"
            elif sigma_val >= 0.6: estado_cav, color_cav = "Riesgo Moderado (Hay Probabilidad)", "orange"
            else:                  estado_cav, color_cav = "¡Alto Riesgo! (Cavitación Crítica)", "red"

            res_map = {
                "Yn": round(y_norm,          2), "A":  round(A_f,               2),
                "V":  round(V_t,             2), "Fr": round(Fr,                2),
                "Winf": round(W_inf_ft * m_conv, 2), "Wsup": round(W_sup_ft * m_conv, 2),
                "W":  round(W_ft * m_conv,   2), "L":  round(1.333 * W_ft * m_conv, 2),
                "f":  round(W_ft / 6 * m_conv, 2), "e": round(W_ft / 12 * m_conv, 2),
                "H":  round(0.75 * W_ft * m_conv, 2), "a": round(0.50 * W_ft * m_conv, 2),
                "b":  round(0.375 * W_ft * m_conv, 2), "c": round(0.50 * W_ft * m_conv, 2),
                "t":  round(W_ft / 12 * m_conv, 2), "Dr": round(W_ft / 20 * m_conv, 2),
                "σ":  round(sigma_val, 2)
            }

            for k, v in res_map.items():
                self.results[k].configure(text=f"{v:.2f}")

            self.status_labels["Cav"].configure(text=estado_cav, text_color=color_cav)

            self.datos_para_exportar = {
                "inputs": {
                    "Tipo Sección": "Rectangular" if t_sec == 1 else "Circular",
                    "Q": Q, "n": n_m, "S": S_m,
                    "b": b_s, "D": D_d, "T": temp_agua, "A": altitud},
                "results": res_map, "cav_diag": estado_cav
            }

            # Advertencia normativa al final — sin bloquear ni borrar resultados
            avisos_vi = []
            if not v_cumple:
                avisos_vi.append(
                    f"• Velocidad (V = {V_t:.2f} m/s) supera el límite de 15.24 m/s (FEMA, 2010)."
                )
            if not q_cumple:
                avisos_vi.append(
                    f"• Caudal (Q = {Q:.2f} m³/s) supera el límite de 11.33 m³/s (FEMA, 2010)."
                )
            if not fr_cumple:
                avisos_vi.append(
                    f"• Froude (Fr = {Fr:.2f}) fuera del rango 1.1 – 10.0 (FEMA, 2010)."
                )
            if avisos_vi:
                messagebox.showwarning(
                    "Advertencia — Límites Normativos USBR VI",
                    "Los resultados se han calculado y se muestran, pero se detectaron las "
                    "siguientes condiciones fuera de norma:\n\n"
                    + "\n".join(avisos_vi)
                    + "\n\nRevise el diseño antes de usar la memoria de cálculo exportada."
                )

        except Exception as e:
            messagebox.showerror("Error de Datos",
                                 f"Verifique los campos numéricos de entrada.\nDetalle: {e}")

    def limpiar(self):
        for e in self.inputs.values():
            old = e.cget("state"); e.configure(state="normal"); e.delete(0, tk.END); e.configure(state=old)
        for r in self.results.values(): r.configure(text="0.00")
        for c in self.checks.values():  c.configure(text="---", text_color="gray")
        self.status_labels["Cav"].configure(text="Evaluar diseño numérico", text_color="gray")
        self.datos_para_exportar = None

    def exportar(self):
        if not self.datos_para_exportar:
            messagebox.showwarning("Atención", "Primero realice un cálculo.")
            return
        path = filedialog.asksaveasfilename(defaultextension=".xlsx",
                                            filetypes=[("Archivos Excel", "*.xlsx")])
        if not path: return
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Resultados USBR VI"
            ws.views.sheetView[0].showGridLines = True

            fill_header  = PatternFill(start_color="1A374D", end_color="1A374D", fill_type="solid")
            fill_section = PatternFill(start_color="406882", end_color="406882", fill_type="solid")
            fill_zebra   = PatternFill(start_color="F9FBFC", end_color="F9FBFC", fill_type="solid")

            font_title   = Font(name=FUENTE_BASE, size=14, bold=True, color="1A374D")
            font_section = Font(name=FUENTE_BASE, size=11, bold=True, color="FFFFFF")
            font_header  = Font(name=FUENTE_BASE, size=10, bold=True, color="FFFFFF")
            font_bold    = Font(name=FUENTE_BASE, size=10, bold=True)
            font_regular = Font(name=FUENTE_BASE, size=10)

            thin = Border(
                left=Side(style='thin', color='B1D0E0'), right=Side(style='thin', color='B1D0E0'),
                top=Side(style='thin', color='B1D0E0'), bottom=Side(style='thin', color='B1D0E0'))

            ws["A1"] = "MEMORIA TÉCNICA DE DISEÑO HIDRÁULICO - USBR TIPO VI"
            ws["A1"].font = font_title

            ws["A4"] = "1. DATOS INICIALES DE ENTRADA"
            ws.merge_cells("A4:D4"); ws["A4"].fill = fill_section; ws["A4"].font = font_section

            headers = ["Parámetro Técnico", "Símbolo", "Valor Numérico", "Unidad"]
            for c_idx, h in enumerate(headers, start=1):
                cell = ws.cell(row=5, column=c_idx, value=h)
                cell.fill = fill_header; cell.font = font_header
                cell.alignment = Alignment(horizontal="center")

            inp = self.datos_para_exportar["inputs"]
            datos_in = [
                ("Geometría de Conducción", "Sección", inp["Tipo Sección"], "-"),
                ("Caudal Máximo de Diseño",  "Q",       inp["Q"],           "m³/s"),
                ("Coeficiente de Manning",   "n",       inp["n"],           "-"),
                ("Pendiente de Canal",       "S",       inp["S"],           "m/m"),
                ("Ancho de Solera",          "b",       inp["b"],           "m"),
                ("Diámetro de Conducción",   "D",       inp["D"],           "m"),
                ("Temperatura del fluido",   "T",       inp["T"],           "°C"),
                ("Altitud del Proyecto",     "A",       inp["A"],           "msnm"),
            ]

            r_curr = 6
            for item in datos_in:
                ws.cell(row=r_curr, column=1, value=item[0]).font = font_regular
                ws.cell(row=r_curr, column=2, value=item[1]).font = font_bold
                v_cell = ws.cell(row=r_curr, column=3, value=item[2]); v_cell.font = font_regular
                if isinstance(item[2], float): v_cell.number_format = "0.00"
                ws.cell(row=r_curr, column=4, value=item[3]).font = font_regular
                for col in range(1, 5): ws.cell(row=r_curr, column=col).border = thin
                r_curr += 1

            r_curr += 1
            ws.cell(row=r_curr, column=1,
                    value="2. DIMENSIONAMIENTO GEOMÉTRICO Y SEGURIDAD").font = font_section
            ws.merge_cells(start_row=r_curr, start_column=1, end_row=r_curr, end_column=4)
            ws.cell(row=r_curr, column=1).fill = fill_section

            r_curr += 1
            for c_idx, h in enumerate(headers, start=1):
                cell = ws.cell(row=r_curr, column=c_idx, value=h)
                cell.fill = fill_header; cell.font = font_header
                cell.alignment = Alignment(horizontal="center")

            r_curr += 1
            res = self.datos_para_exportar["results"]
            datos_out = [
                ("Tirante Normal",           "Yn",    res["Yn"],   "m"),
                ("Área Mojada",              "A",     res["A"],    "m²"),
                ("Velocidad de Ingreso",     "V",     res["V"],    "m/s"),
                ("Número de Froude",         "Fr",    res["Fr"],   "-"),
                ("Ancho Mínimo (FEMA)",      "W_inf", res["Winf"], "m"),
                ("Ancho Máximo (FEMA)",      "W_sup", res["Wsup"], "m"),
                ("Ancho de Diseño",          "W",     res["W"],    "m"),
                ("Longitud del Cuenco",      "L",     res["L"],    "m"),
                ("Espesor Frontal",          "f",     res["f"],    "m"),
                ("Transición de Salida",     "e",     res["e"],    "m"),
                ("Altura Total del Cuenco",  "H",     res["H"],    "m"),
                ("Lado del Bloque",          "a",     res["a"],    "m"),
                ("Ancho del Deflector",      "b",     res["b"],    "m"),
                ("Contrafuerte",             "c",     res["c"],    "m"),
                ("Espesor Mínimo Pantalla",  "t",     res["t"],    "m"),
                ("Diámetro de Roca",         "Dr",    res["Dr"],   "m"),
                ("Índice de Cavitación",     "σ",     res["σ"],    "-"),
                ("Diagnóstico Cavitación", "Estado", self.datos_para_exportar["cav_diag"], "—"),
            ]

            for item in datos_out:
                ws.cell(row=r_curr, column=1, value=item[0]).font = font_regular
                ws.cell(row=r_curr, column=2, value=item[1]).font = font_bold
                v_cell = ws.cell(row=r_curr, column=3, value=item[2])
                if isinstance(item[2], float):
                    v_cell.font = font_regular; v_cell.number_format = "0.00"
                else:
                    v_cell.font = font_bold
                ws.cell(row=r_curr, column=4, value=item[3]).font = font_regular
                if r_curr % 2 == 0:
                    for col in range(1, 5): ws.cell(row=r_curr, column=col).fill = fill_zebra
                for col in range(1, 5): ws.cell(row=r_curr, column=col).border = thin
                r_curr += 1

            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                ws.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 3, 11)
            ws.column_dimensions['A'].width = 36
            ws.column_dimensions['C'].width = 45

            #  Esquemas técnicos USBR VI 
            r_curr += 2
            ws.cell(row=r_curr, column=1,
                    value="3. ESQUEMAS TÉCNICOS DEL CUENCO USBR TIPO VI").font = font_section
            ws.merge_cells(start_row=r_curr, start_column=1, end_row=r_curr, end_column=4)
            ws.cell(row=r_curr, column=1).fill = fill_section
            r_curr += 2

            _esquemas_usbr = [
                ("Vista en Planta USBR VI",  resolver_ruta("Imagenes/Vista en planta_USBR_VI.png")),
                ("Sección Transversal A-A",  resolver_ruta("Imagenes/Sección A-A_USBR_VI.png")),
            ]
            from PIL import Image as _PILImg
            for etiqueta, ruta_img in _esquemas_usbr:
                if os.path.exists(ruta_img):
                    ws.cell(row=r_curr, column=1, value=etiqueta).font = font_bold
                    r_curr += 1
                    try:
                        _pi = _PILImg.open(ruta_img)
                        _ow, _oh = _pi.size
                        _tw = 550
                        _th = max(60, int(_oh * _tw / _ow))
                        _xi = OXLImage(ruta_img)
                        _xi.width  = _tw
                        _xi.height = _th
                        ws.add_image(_xi, f"A{r_curr}")
                        _rows = (_th // 18) + 2
                        for _rr in range(r_curr, r_curr + _rows):
                            ws.row_dimensions[_rr].height = _th / _rows
                        r_curr += _rows
                    except Exception:
                        r_curr += 3
                else:
                    r_curr += 1

            wb.save(path)
            messagebox.showinfo("Éxito", "Memoria técnica exportada correctamente a Excel.")
        except PermissionError:
            messagebox.showerror("Error de Archivo",
                                 "El archivo Excel está abierto. Ciérrelo e intente de nuevo.")
        except Exception as e:
            messagebox.showerror("Error de Archivo",
                                 f"No se pudo guardar el archivo.\nDetalle: {e}")

# ------------------------------------------------------------------------------
# 8. CONTROLADOR DE NAVEGACIÓN
# ------------------------------------------------------------------------------
class AplicacionPrincipal(ctk.CTk):
    ANCHO, ALTO = 1450, 920

    def __init__(self):
        super().__init__()
        self.title("Hidra Dis - Cuencos Disipadores")
        self.geometry(f"{self.ANCHO}x{self.ALTO}")
        self.minsize(1250, 780)
        centrar_ventana(self, self.ANCHO, self.ALTO)

        try:
            self.iconbitmap(resolver_ruta("icono_Hidra_Dis.ico"))
        except Exception:
            pass

        self.grid_rowconfigure(0, weight=1)
        self.grid_columnconfigure(0, weight=1)

        self.pantalla_inicio  = PantallaInicio(master=self,  control_navegacion=self,
                                                ancho_ref=self.ANCHO, alto_ref=self.ALTO)
        self.pantalla_saf     = InterfaceSAF(master=self,    control_navegacion=self)
        self.pantalla_usbr_vi = InterfaceUSBRVI(master=self, control_navegacion=self)

        self.mostrar_inicio()

    def ocultar_vistas(self):
        self.pantalla_inicio.grid_forget()
        self.pantalla_saf.grid_forget()
    def ocultar_vistas(self):
        self.pantalla_inicio.grid_forget()
        self.pantalla_saf.grid_forget()
        self.pantalla_usbr_vi.grid_forget()

    def mostrar_inicio(self):
        self.ocultar_vistas()
        self.pantalla_inicio.grid(row=0, column=0, sticky="nsew")

    def mostrar_saf(self):
        self.ocultar_vistas()
        self.pantalla_saf.grid(row=0, column=0, sticky="nsew")

    def mostrar_usbr_vi(self):
        self.ocultar_vistas()
        self.pantalla_usbr_vi.grid(row=0, column=0, sticky="nsew")

if __name__ == "__main__":
    app = AplicacionPrincipal()
    app.mainloop()
