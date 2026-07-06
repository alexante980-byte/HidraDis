import os
import sys
import numpy as np
import openpyxl
from datetime import datetime
import customtkinter as ctk
import tkinter as tk
from tkinter import filedialog, messagebox
from PIL import Image
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# Configuración general de apariencia
ctk.set_appearance_mode("Dark")
ctk.set_default_color_theme("blue")

# ------------------------------------------------------------------------------
# EXCEPCIONES PERSONALIZADAS DE DISEÑO NORMADO
# ------------------------------------------------------------------------------
class SAFValidationError(Exception):
    """Excepción para detener el cálculo si se violan los límites hidráulicos del cuenco SAF."""
    pass

# ------------------------------------------------------------------------------
# MOTOR HIDRÁULICO 
# ------------------------------------------------------------------------------
def resolver_ruta(ruta_relativa):
    """Gestiona las rutas de recursos tanto en desarrollo como al compilar con PyInstaller (.exe)."""
    try:
        base_path = sys._MEIPASS
    except AttributeError:
        base_path = os.path.abspath(".")
    return os.path.join(base_path, ruta_relativa)

def calcular_presion_vapor(temperatura_c: float) -> float:
    """Calcula la presión de vapor del agua (Hv) en m.c.a. usando Antoine modificada."""
    p_mmhg = 10**(8.07131 - (1730.63 / (temperatura_c + 233.426)))
    return 0.0136 * p_mmhg

def calcular_presion_atmosferica(altitud_msnm: float) -> float:
    """Calcula la presión atmosférica (H_atm) en m.c.a. según la altitud local."""
    return 10.33 - (altitud_msnm / 900.0)

# ------------------------------------------------------------------------------
# VISTA: PANTALLA DE BIENVENIDA
# ------------------------------------------------------------------------------
class PantallaInicio(ctk.CTkFrame):
    def __init__(self, master, control_navegacion, **kwargs):
        super().__init__(master, **kwargs)
        self.configure(fg_color="#0F172A")
        self._img_refs = []

        # ── ENCABEZADO INSTITUCIONAL ──────────────────────────────────────────
        hdr = ctk.CTkFrame(self, fg_color="#1E293B", corner_radius=0)
        hdr.pack(fill="x")

        hl = ctk.CTkFrame(hdr, fg_color="transparent")
        hl.pack(fill="x", padx=35, pady=12)

        try:
            img_izq = ctk.CTkImage(
                light_image=Image.open(resolver_ruta(os.path.join("Imagenes", "Escuela_Politécnica_Nacional.png"))),
                dark_image=Image.open(resolver_ruta(os.path.join("Imagenes", "Escuela_Politécnica_Nacional.png"))),
                size=(68, 68))
            self._img_refs.append(img_izq)
            ctk.CTkLabel(hl, image=img_izq, text="").pack(side="left", padx=(0, 20))
        except Exception:
            pass

        try:
            img_der = ctk.CTkImage(
                light_image=Image.open(resolver_ruta(os.path.join("Imagenes", "logo_fica.png"))),
                dark_image=Image.open(resolver_ruta(os.path.join("Imagenes", "logo_fica.png"))),
                size=(135, 43))
            self._img_refs.append(img_der)
            ctk.CTkLabel(hl, image=img_der, text="").pack(side="right", padx=(20, 0))
        except Exception:
            pass

        tc = ctk.CTkFrame(hl, fg_color="transparent")
        tc.pack(side="left", expand=True)
        ctk.CTkLabel(tc, text="ESCUELA POLITÉCNICA NACIONAL",
                     font=("Segoe UI", 16, "bold"), text_color="#F8FAFC").pack()
        ctk.CTkLabel(tc, text="FACULTAD DE INGENIERÍA CIVIL Y AMBIENTAL",
                     font=("Segoe UI", 12, "bold"), text_color="#38BDF8").pack(pady=1)
        ctk.CTkLabel(tc, text="TRABAJO DE INTEGRACIÓN CURRICULAR",
                     font=("Segoe UI", 10, "italic"), text_color="#94A3B8").pack()

        ctk.CTkFrame(self, fg_color="#0EA5E9", height=4, corner_radius=0).pack(fill="x")

        # ── TÍTULO PRINCIPAL ──────────────────────────────────────────────────
        title_area = ctk.CTkFrame(self, fg_color="transparent")
        title_area.pack(pady=(40, 0))

        ctk.CTkLabel(title_area, text="HIDRA DIS PRO",
                     font=("Segoe UI", 46, "bold"), text_color="#F8FAFC").pack()
        ctk.CTkLabel(title_area,
                     text="Software Integrado de Diseño y Evaluación de Cuencos Disipadores de Energía",
                     font=("Segoe UI", 14, "italic"), text_color="#38BDF8").pack(pady=4)

        ctk.CTkFrame(self, fg_color="#334155", height=1, corner_radius=0).pack(
            fill="x", padx=100, pady=(15, 20))

        # ── DESCRIPCIÓN ───────────────────────────────────────────────────────
        ctk.CTkLabel(self,
                     text=(
                         "Dimensionamiento geométrico automatizado y evaluación frente a fenómenos de cavitación\n"
                         "en obras de descarga hidráulica bajo normativas de diseño estructural y de fluidos.\n"
                         "Módulos validados para estructuras Tipo SAF y USBR Tipo VI."
                     ),
                     font=("Segoe UI", 13), text_color="#E2E8F0",
                     justify="center").pack(pady=(0, 35))

        # ── TARJETAS DE MÓDULO ────────────────────────────────────────────────
        cards_frame = ctk.CTkFrame(self, fg_color="transparent")
        cards_frame.pack(pady=10)

        # Tarjeta SAF
        saf_card = ctk.CTkFrame(cards_frame, fg_color="#1E293B", corner_radius=16,
                                 border_width=1, border_color="#38BDF8",
                                 width=350, height=210)
        saf_card.grid(row=0, column=0, padx=30)
        saf_card.pack_propagate(False)

        badge_saf = ctk.CTkFrame(saf_card, fg_color="#0EA5E9", corner_radius=8,
                                  width=120, height=32)
        badge_saf.pack(pady=(20, 6))
        badge_saf.pack_propagate(False)
        ctk.CTkLabel(badge_saf, text="CUENCO SAF",
                     font=("Segoe UI", 12, "bold"), text_color="#0F172A").pack(expand=True)

        ctk.CTkLabel(saf_card, text="Saint Anthony Falls Laboratory",
                     font=("Segoe UI", 11, "italic"), text_color="#7DD3FC").pack(pady=(0, 4))
        ctk.CTkLabel(saf_card, text="Fr₁: 1.7 – 17  |  Colectores y Azudes ancho variable",
                     font=("Segoe UI", 10), text_color="#94A3B8").pack(pady=(0, 15))

        ctk.CTkButton(saf_card, text="Ir al módulo SAF  →",
                      font=("Segoe UI", 12, "bold"), width=220, height=36,
                      fg_color="#0284C7", hover_color="#0369A1", corner_radius=8,
                      text_color="#FFFFFF",
                      command=control_navegacion.mostrar_saf).pack(pady=4)

        # Tarjeta USBR VI
        vi_card = ctk.CTkFrame(cards_frame, fg_color="#1E293B", corner_radius=16,
                                border_width=1, border_color="#38BDF8",
                                width=350, height=210)
        vi_card.grid(row=0, column=1, padx=30)
        vi_card.pack_propagate(False)

        badge_vi = ctk.CTkFrame(vi_card, fg_color="#0EA5E9", corner_radius=8,
                                 width=140, height=32)
        badge_vi.pack(pady=(20, 6))
        badge_vi.pack_propagate(False)
        ctk.CTkLabel(badge_vi, text="CUENCO USBR VI",
                     font=("Segoe UI", 12, "bold"), text_color="#0F172A").pack(expand=True)

        ctk.CTkLabel(vi_card, text="Bureau of Reclamation — Tipo VI",
                     font=("Segoe UI", 11, "italic"), text_color="#7DD3FC").pack(pady=(0, 4))
        ctk.CTkLabel(vi_card, text="V ≤ 15.24 m/s  |  Q ≤ 11.33 m³/s (Impacto)",
                     font=("Segoe UI", 10), text_color="#94A3B8").pack(pady=(0, 15))

        ctk.CTkButton(vi_card, text="Ir al módulo USBR VI  →",
                      font=("Segoe UI", 12, "bold"), width=220, height=36,
                      fg_color="#0284C7", hover_color="#0369A1", corner_radius=8,
                      text_color="#FFFFFF",
                      command=control_navegacion.mostrar_usbr_vi).pack(pady=4)

        # ── PIE DE PANTALLA ───────────────────────────────────────────────────
        footer_ini = ctk.CTkFrame(self, fg_color="#0F172A", corner_radius=0, height=40)
        footer_ini.pack(fill="x", side="bottom")
        footer_ini.pack_propagate(False)
        ctk.CTkLabel(footer_ini,
                     text=f"Autor: Alex Ante   |   Ingeniería Civil — EPN   |   {datetime.now().strftime('%d/%m/%Y')}",
                     font=("Segoe UI", 10), text_color="#64748B").pack(expand=True)

# ------------------------------------------------------------------------------
# VISTA: MODULO COMPLETO - CUENCO TIPO SAF
# ------------------------------------------------------------------------------
class InterfaceSAF(ctk.CTkFrame):
    def __init__(self, master, control_navegacion, **kwargs):
        super().__init__(master, **kwargs)
        self.configure(fg_color="#0F172A")
        self._img_refs = []
        self.inputs = {}
        self.input_frames = {}
        self.results = {}
        self.status_labels = {}
        self.data_export = {}
        self.control_navegacion = control_navegacion

        self.ruta_epn        = resolver_ruta(os.path.join("Imagenes", "Escuela_Politécnica_Nacional.png"))
        self.ruta_fica       = resolver_ruta(os.path.join("Imagenes", "logo_fica.png"))
        self.ruta_planta     = resolver_ruta(os.path.join("Imagenes", "Vista en planta SAF.png"))
        self.ruta_elevacion  = resolver_ruta(os.path.join("Imagenes", "Vista en elevación_SAF.png"))
        self.ruta_isometrica = resolver_ruta(os.path.join("Imagenes", "Vista isométrica_SAF.png"))

        self.init_module_ui()

    def init_module_ui(self):
        header = ctk.CTkFrame(self, fg_color="#1E293B", corner_radius=0)
        header.pack(fill="x", padx=0, pady=0)
        content_head = ctk.CTkFrame(header, fg_color="transparent")
        content_head.pack(pady=8, fill="x")

        self.btn_volver = ctk.CTkButton(
            content_head, text="⬅ Menú Principal", width=140, height=32,
            fg_color="#334155", hover_color="#475569", font=("Segoe UI", 11, "bold"),
            command=self.control_navegacion.mostrar_inicio)
        self.btn_volver.pack(side="left", padx=20)

        self.load_img(content_head, self.ruta_epn, (45, 45), "left")
        titles_frame = ctk.CTkFrame(content_head, fg_color="transparent")
        titles_frame.pack(side="left", padx=15, expand=True)
        ctk.CTkLabel(titles_frame, text="DISEÑO DE CUENCO DISIPADOR TIPO SAF",
                     font=("Segoe UI", 18, "bold"), text_color="#F8FAFC").pack()
        ctk.CTkLabel(titles_frame,
                     text="Cálculo hidráulico y dimensionamiento geométrico normado por el Saint Anthony Falls Laboratory.",
                     font=("Segoe UI", 11, "italic"), text_color="#94A3B8").pack()
        self.load_img(content_head, self.ruta_fica, (110, 35), "right")

        body = ctk.CTkFrame(self, fg_color="transparent")
        body.pack(fill="both", expand=True, padx=15, pady=10)

        type_frame = ctk.CTkFrame(body, fg_color="#1E293B", corner_radius=8, border_width=1, border_color="#334155")
        type_frame.pack(fill="x", pady=(0, 8))
        ctk.CTkLabel(type_frame, text="Tipo de Configuración de Cuenco:",
                     font=("Segoe UI", 12, "bold"), text_color="#38BDF8").pack(side="left", padx=15, pady=6)
        
        self.tipo_var = ctk.StringVar(value="Colector")
        ctk.CTkRadioButton(type_frame, text="Azud (Ancho Constante)",
                           variable=self.tipo_var, value="Derivacion",
                           command=self.toggle_inputs, font=("Segoe UI", 11)).pack(side="left", padx=15)
        ctk.CTkRadioButton(type_frame, text="Colector (Ancho Variable)",
                           variable=self.tipo_var, value="Colector",
                           command=self.toggle_inputs, font=("Segoe UI", 11)).pack(side="left", padx=15)

        panels = ctk.CTkFrame(body, fg_color="transparent")
        panels.pack(fill="both", expand=True)

        left_side = ctk.CTkFrame(panels, fg_color="transparent", width=680)
        left_side.pack(side="left", fill="both", expand=False, padx=(0, 10))
        left_side.pack_propagate(False)
        self.left_scroll = ctk.CTkScrollableFrame(left_side, fg_color="#1E293B", label_text="")
        self.left_scroll.pack(fill="both", expand=True)

        right_side = ctk.CTkFrame(panels, fg_color="transparent", width=710)
        right_side.pack(side="right", fill="both", expand=True)
        right_side.pack_propagate(False)
        self.right_scroll = ctk.CTkScrollableFrame(right_side, fg_color="#1E293B")
        self.right_scroll.pack(fill="both", expand=True)

        self.setup_sections()
        self.setup_graphics()

        # ── PIE DE PÁGINA ─────────────────────────────────────────────────────
        footer = ctk.CTkFrame(self, fg_color="#1E293B", height=80, corner_radius=0, border_width=1, border_color="#334155")
        footer.pack(fill="x", side="bottom")
        footer.pack_propagate(False)

        info_tic = ctk.CTkFrame(footer, fg_color="transparent")
        info_tic.pack(side="left", padx=25, pady=5)
        ctk.CTkLabel(info_tic, text="Facultad de Ingeniería Civil y Ambiental — EPN",
                     font=("Segoe UI", 11, "bold"), text_color="#F8FAFC").pack(anchor="w")
        ctk.CTkLabel(info_tic, text="Trabajo de Integración Curricular  |  Disipador SAF",
                     font=("Segoe UI", 10), text_color="#94A3B8").pack(anchor="w")

        btn_frame = ctk.CTkFrame(footer, fg_color="transparent")
        btn_frame.pack(side="right", padx=25, pady=20)
        ctk.CTkButton(btn_frame, text="Calcular", command=self.calcular,
                      fg_color="#0EA5E9", hover_color="#0284C7", text_color="#0F172A", width=100, height=32,
                      font=("Segoe UI", 12, "bold")).pack(side="left", padx=5)
        ctk.CTkButton(btn_frame, text="Limpiar", command=self.limpiar,
                      fg_color="#475569", hover_color="#334155", text_color="#F8FAFC", width=90, height=32,
                      font=("Segoe UI", 11)).pack(side="left", padx=5)
        ctk.CTkButton(btn_frame, text="Exportar Excel", command=self.exportar_excel,
                      fg_color="#16A34A", hover_color="#15803D", text_color="#FFFFFF", width=130, height=32,
                      font=("Segoe UI", 11, "bold")).pack(side="left", padx=5)

    def load_img(self, parent, path, size, side):
        if os.path.exists(path):
            try:
                img = ctk.CTkImage(light_image=Image.open(path), dark_image=Image.open(path), size=size)
                self._img_refs.append(img)
                ctk.CTkLabel(parent, image=img, text="").pack(side=side, padx=5)
            except Exception:
                pass

    def toggle_inputs(self):
        if self.tipo_var.get() == "Derivacion":
            self.input_frames["D0_frame"].grid_forget()
            self.input_frames["Br_frame"].grid(row=0, column=1, padx=6, pady=4)
        else:
            self.input_frames["Br_frame"].grid_forget()
            self.input_frames["D0_frame"].grid(row=0, column=1, padx=6, pady=4)

    def setup_sections(self):
        ctk.CTkLabel(self.left_scroll, text="Datos de Entrada:",
                     font=("Segoe UI", 14, "bold"), text_color="#38BDF8",
                     anchor="w").pack(fill="x", padx=10, pady=(5, 2))
        in_f = ctk.CTkFrame(self.left_scroll, fg_color="transparent")
        in_f.pack(fill="x", pady=(1, 10))

        grid = ctk.CTkFrame(in_f, fg_color="transparent")
        grid.pack(padx=5, pady=4, anchor="w")

        self.input_frames["Q_frame"]    = self.create_input(grid, "Caudal (Q):",              "m³/s", "Q",   row=0, col=0)
        self.input_frames["Br_frame"]   = self.create_input(grid, "Ancho azud (Bᵣ):",         "m",    "Br",  row=0, col=1)
        self.input_frames["D0_frame"]   = self.create_input(grid, "Diámetro colector (D₀):",  "m",    "D0",  row=0, col=1)
        self.input_frames["y1_frame"]   = self.create_input(grid, "Tirante supercrítico (y₁):", "m",   "y1",  row=1, col=0)
        self.input_frames["Temp_frame"] = self.create_input(grid, "Temperatura (T):",          "°C",   "Temp",row=1, col=1)
        self.input_frames["Alt_frame"]  = self.create_input(grid, "Altitud (A):",              "msnm", "Alt", row=2, col=0)

        self.toggle_inputs()

        ctk.CTkLabel(self.left_scroll, text="Resultados del Dimensionamiento:",
                     font=("Segoe UI", 14, "bold"), text_color="#38BDF8",
                     anchor="w").pack(fill="x", padx=10, pady=(5, 2))
        res_f = ctk.CTkFrame(self.left_scroll, fg_color="transparent")
        res_f.pack(fill="both", expand=True, pady=(1, 5))

        sections = [
            ("1. Parámetros de Cuenco", [
                ("Ancho Cuenco", "Wb", "Wb", "m"),
                ("Ancho Rápida", "Wb₁", "Wb1", "m"),
                ("Ancho Intermedio", "Wb₂", "Wb2", "m")]),
            ("2. Parámetros Hidráulicos Reales", [
                ("Número Froude", "Fr₁", "Fr1_ap", "-"),
                ("Velocidad Flujo", "V₁", "V1", "m/s"),
                ("Tirante Crítico", "y꜀", "yc", "m")]),
            ("3. Condiciones del cuenco SAF", [
                ("Tirante Conjugado", "y₂", "d2", "m"),
                ("Factor Corrección", "C", "C", "-")]),
            ("4. Bloques de Rápida", [
                ("Número Bloques", "n꜀♭", "ncb", "u"),
                ("Ancho", "w꜀♭", "wcb", "m"),
                ("Altura", "h꜀♭", "hcb", "m")]),
            ("5. Bloques de Fondo", [
                ("Número Bloques", "nբ♭", "nfb", "u"),
                ("Ancho", "wբ♭", "wfb", "m"),
                ("Altura", "hբ♭", "hfb", "m"),
                ("Espesor Cresta", "e", "e", "m"),
                ("Distancia Pared", "dᵣₑₐₗ", "dreal", "m")]),
            ("6. Aproximación Geométrica", [
                ("Transición Ent.", "Lₐ", "LA", "m"),
                ("Longitud Cuenco", "LB", "LB", "m"),
                ("Ubicación Bloques", "Dₑₛ", "Des", "m")]),
            ("7. Estructura Final", [
                ("Altura Umbral", "h₄", "h4", "m"),
                ("Altura Pared", "H₆", "h6", "m"),
                ("Borde Libre", "z", "z_bl", "m"),
                ("Ancho Salida", "Wb₃", "Wb3", "m")]),
            ("8. Cavitación (Ingreso de la rápida)", [
                ("Índice de cavitación", "σ", "sigma", "-")])
        ]

        for title, items in sections:
            self.create_result_group(res_f, title, items)

    def create_input(self, parent, label, unit, key, row, col):
        f = ctk.CTkFrame(parent, fg_color="transparent")
        f.grid(row=row, column=col, padx=12, pady=5, sticky="w")
        ctk.CTkLabel(f, text=label, width=140, font=("Segoe UI", 11), anchor="w", text_color="#E2E8F0").pack(side="left")
        e = ctk.CTkEntry(f, width=75, height=24, fg_color="#0F172A", text_color="#F8FAFC", border_color="#334155")
        e.pack(side="left", padx=2)
        ctk.CTkLabel(f, text=unit, font=("Segoe UI", 10), text_color="#94A3B8",
                     width=40, anchor="w").pack(side="left", padx=2)
        self.inputs[key] = e
        return f

    def create_result_group(self, parent, title, items):
        ctk.CTkLabel(parent, text=title, font=("Segoe UI", 12, "bold"),
                     text_color="#7DD3FC", anchor="w").pack(fill="x", padx=12, pady=(8, 2))
        container = ctk.CTkFrame(parent, fg_color="transparent")
        container.pack(fill="x", padx=8, pady=2)

        for col_idx in range(3):
            container.grid_columnconfigure(col_idx, weight=1, minsize=215)

        for i, item in enumerate(items):
            r, c = divmod(i, 3)
            cell = ctk.CTkFrame(container, fg_color="transparent")
            cell.grid(row=r, column=c, padx=4, pady=3, sticky="w")

            ctk.CTkLabel(cell, text=f"{item[0]} [{item[1]}]:",
                         width=110, anchor="w", font=("Segoe UI", 10), text_color="#CBD5E1").pack(side="left")
            res_box = ctk.CTkLabel(cell, text="0.00", width=70, height=22,
                                    fg_color="#0F172A", text_color="#38BDF8", corner_radius=4,
                                    font=("Segoe UI", 10, "bold"))
            res_box.pack(side="left", padx=2)
            ctk.CTkLabel(cell, text=item[3], width=35, text_color="#64748B",
                         font=("Segoe UI", 9.5), anchor="w").pack(side="left")

        if "SAF" in title:
            sf = ctk.CTkFrame(container, fg_color="transparent")
            sf.grid(row=2, column=0, columnspan=3, pady=4, sticky="w")
            ctk.CTkLabel(sf, text="Verificación (1.7 < Fr₁ < 17):", font=("Segoe UI", 10), text_color="#94A3B8").pack(side="left", padx=(2, 0))
            self.status_labels["Fr1"] = ctk.CTkLabel(sf, text="--", font=("Segoe UI", 10, "bold"), text_color="#64748B")
            self.status_labels["Fr1"].pack(side="left", padx=6)

        elif "Fondo" in title:
            sf = ctk.CTkFrame(container, fg_color="transparent")
            sf.grid(row=2, column=0, columnspan=3, sticky="w", pady=4)
            ctk.CTkLabel(sf, text="Ocupación (40-55% Bᵣ):", font=("Segoe UI", 10), text_color="#94A3B8").pack(side="left", padx=(2, 0))
            self.results["Oc"] = ctk.CTkLabel(sf, text="0.0%", width=55, font=("Segoe UI", 10, "bold"), text_color="#38BDF8")
            self.results["Oc"].pack(side="left", padx=4)
            self.status_labels["Oc"] = ctk.CTkLabel(sf, text="--", font=("Segoe UI", 10, "bold"), text_color="#64748B")
            self.status_labels["Oc"].pack(side="left", padx=4)

        elif "Cavitación" in title:
            sf = ctk.CTkFrame(container, fg_color="transparent")
            sf.grid(row=1, column=0, columnspan=3, sticky="w", pady=4)
            ctk.CTkLabel(sf, text="Diagnóstico Cavitación:", font=("Segoe UI", 10, "bold"), text_color="#F8FAFC").pack(side="left", padx=(2, 0))
            self.status_labels["Cav"] = ctk.CTkLabel(sf, text="Evaluar diseño", font=("Segoe UI", 10, "bold"), text_color="#64748B")
            self.status_labels["Cav"].pack(side="left", padx=6)

    def setup_graphics(self):
        ctk.CTkLabel(self.right_scroll, text="Esquema Geométrico Estructural SAF:",
                     font=("Segoe UI", 14, "bold"), text_color="#38BDF8",
                     anchor="w").pack(fill="x", padx=10, pady=(5, 4))

        formatos = {
            "Vista en Planta":    (520, 310),
            "Vista en Elevación": (520, 310),
            "Vista Isométrica":   (520, 310)
        }

        for title, path in [
            ("Vista en Planta",    self.ruta_planta),
            ("Vista en Elevación", self.ruta_elevacion),
            ("Vista Isométrica",   self.ruta_isometrica)
        ]:
            f = ctk.CTkFrame(self.right_scroll, fg_color="#1E293B", corner_radius=8, border_width=1, border_color="#334155")
            f.pack(fill="x", pady=5, padx=5)
            ctk.CTkLabel(f, text=title, font=("Segoe UI", 11, "bold"), text_color="#F8FAFC").pack(pady=4)

            if os.path.exists(path):
                w, h = formatos[title]
                img = ctk.CTkImage(light_image=Image.open(path), dark_image=Image.open(path), size=(w, h))
                self._img_refs.append(img)
                ctk.CTkLabel(f, image=img, text="").pack(pady=4, padx=12)
            else:
                ctk.CTkLabel(f, text=f"[{title} — Esquema de referencia técnica]",
                             font=("Segoe UI", 10, "italic"), text_color="#64748B").pack(pady=25)

            ctk.CTkLabel(f, text="Adaptado de Chow (1994)", font=("Segoe UI", 9, "italic"), text_color="#64748B").pack(pady=(0, 4), anchor="e", padx=12)

    def calcular(self):
        try:
            g = 9.81
            mode = self.tipo_var.get()

            Q        = float(self.inputs["Q"].get().replace(',', '.'))
            y1       = float(self.inputs["y1"].get().replace(',', '.'))
            temp_agua= float(self.inputs["Temp"].get().replace(',', '.'))
            altitud  = float(self.inputs["Alt"].get().replace(',', '.'))

            if mode == "Derivacion":
                Br  = float(self.inputs["Br"].get().replace(',', '.'))
                V1  = Q / (Br * y1)
                Wb1 = Br
                Wb  = Br
                LA  = 0.0
            else:
                D0  = float(self.inputs["D0"].get().replace(',', '.'))
                Wb1 = D0 if (Q / (D0 * y1)) < 6 else 2.5 * D0 if (Q / (D0 * y1)) < 12 else 3.0 * D0
                V1  = Q / (Wb1 * y1)
                Wb  = max(D0, 1.70 * Q / (np.sqrt(g) * D0**1.5))
                LA  = (Wb1 - D0) / 2.0

            Fr1 = V1 / np.sqrt(g * y1)

            if Fr1 < 1.7 or Fr1 > 17.0:
                raise SAFValidationError(
                    f"El número de Froude calculado (Fr₁ = {Fr1:.2f}) se encuentra fuera del rango empírico normado "
                    f"para cuencos SAF (Límites: 1.7 ≤ Fr₁ ≤ 17).\n\n"
                    f"Ajuste los parámetros geométricos o el caudal de entrada."
                )

            yc  = ((Q / Wb)**2 / g)**(1/3)
            d2  = (y1 / 2) * (np.sqrt(1 + 8 * Fr1**2) - 1)

            if   1.7 < Fr1 <= 5.5:  C_val = 1.1 - (Fr1**2 / 120)
            elif 5.5 < Fr1 <= 11:   C_val = 0.85
            else:                    C_val = 1.0 - (Fr1**2 / 800)

            LB  = (4.5 * d2) / (C_val * (Fr1**0.76))
            Wb2 = Wb1 + (Wb - Wb1) / 3.0 if mode == "Colector" else Wb

            ncb  = int(Wb1 / (1.5 * y1)) if (1.5 * y1) > 0 else 0
            wcb  = Wb1 / (2 * ncb) if ncb > 0 else 0
            nfb  = ncb
            wfb  = Wb2 / (2 * nfb) if nfb > 0 else 0

            Hatm      = calcular_presion_atmosferica(altitud)
            Hv        = calcular_presion_vapor(temp_agua)
            carga_vel = (V1**2) / (2 * g)
            sigma     = (Hatm - Hv - y1) / carga_vel if carga_vel > 0 else 0

            if   sigma > 1.0:        estado_cav, color_cav = "NO PRESENTARÁ CAVITACIÓN ✔", "#4ADE80"
            elif 0.4 <= sigma <= 1.0: estado_cav, color_cav = "RIESGO MODERADO DE CAVITACIÓN ⚠", "#FBBF24"
            else:                    estado_cav, color_cav = "ALTO RIESGO DE CAVITACIÓN ✘", "#EF4444"

            res_map = {
                "Wb": Wb, "Wb1": Wb1, "Wb2": Wb2, "yc": yc, "V1": V1,
                "Fr1_ap": Fr1, "Fr1": Fr1, "d2": d2, "C": C_val,
                "LA": LA, "LB": LB, "Des": LB / 3.0,
                "ncb": ncb, "wcb": wcb, "acb": wcb, "hcb": y1,
                "nfb": nfb, "wfb": wfb, "afb": wfb, "hfb": y1,
                "e": 0.5 * y1,
                "dreal": (Wb2 - (nfb * wfb) - (nfb - 1) * wfb) / 2 if nfb > 0 else 0,
                "h4": max(1.0, 0.07 * (d2 / C_val)),
                "h6": d2 * (1.0 + (1.0 / (3.0 * C_val))),
                "z_bl": d2 / 3.0, "Wb3": Wb + 2.0 * LB, "sigma": sigma
            }

            for k, v in res_map.items():
                if k in self.results:
                    self.results[k].configure(
                        text=f"{v:.4f}" if k == "sigma"
                        else f"{v:.3f}" if k in ["Fr1_ap", "yc", "V1", "C"]
                        else f"{v:.2f}")

            ocup = (nfb * wfb / Wb2) * 100 if nfb > 0 else 0
            self.results["Oc"].configure(text=f"{ocup:.1f}%")
            self.status_labels["Fr1"].configure(text="CUMPLE NORMADO", text_color="#4ADE80")
            self.status_labels["Oc"].configure(
                text="CUMPLE (40-55%)" if 40 <= ocup <= 55 else "REVISAR OCUPACIÓN",
                text_color="#4ADE80" if 40 <= ocup <= 55 else "#EF4444")
            self.status_labels["Cav"].configure(text=estado_cav, text_color=color_cav)

            self.data_export = {
                "inputs": {
                    "Tipo Proyecto": mode, "Q": Q, "y1": y1,
                    "Temp": temp_agua, "Alt": altitud,
                    "Br": Br if mode == "Derivacion" else 0.0,
                    "D0": D0 if mode == "Colector" else 0.0},
                "results": res_map, "cav_text": estado_cav
            }

        except SAFValidationError as val_err:
            self.limpiar()
            messagebox.showwarning("Límite Hidráulico Excedido", str(val_err))
        except Exception as err:
            messagebox.showerror("Error de Consistencia", f"Verifique que todos los campos sean numéricos válidos.\nDetalle: {err}")

    def limpiar(self):
        for r in self.results.values(): r.configure(text="0.00")
        for s in self.status_labels.values(): s.configure(text="--", text_color="#64748B")
        if "Oc" in self.results: self.results["Oc"].configure(text="0.0%")
        if "Cav" in self.status_labels:
            self.status_labels["Cav"].configure(text="Evaluar diseño", text_color="#64748B")
        self.data_export = {}

    def exportar_excel(self):
        if not self.data_export:
            messagebox.showwarning("Atención", "Por favor, realice el cálculo antes de exportar.")
            return
        path = filedialog.asksaveasfilename(defaultextension=".xlsx",
                                            filetypes=[("Archivos Excel", "*.xlsx")])
        if not path: return
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Memoria SAF"
            ws.views.sheetView[0].showGridLines = True

            fill_header  = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            fill_section = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
            fill_zebra   = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

            font_title   = Font(name="Segoe UI", size=15, bold=True, color="1E293B")
            font_section = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            font_header  = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
            font_bold    = Font(name="Segoe UI", size=10, bold=True)
            font_regular = Font(name="Segoe UI", size=10)

            thin = Border(
                left=Side(style='thin', color='CBD5E1'), right=Side(style='thin', color='CBD5E1'),
                top=Side(style='thin', color='CBD5E1'), bottom=Side(style='thin', color='CBD5E1'))

            ws["A1"] = "MEMORIA TÉCNICA DE DISEÑO HIDRÁULICO - CUENCO SAF"
            ws["A1"].font = font_title

            ws["A4"] = "1. DATOS DE ENTRADA"
            ws.merge_cells("A4:D4")
            ws["A4"].fill = fill_section
            ws["A4"].font = font_section

            headers = ["Parámetro", "Variable", "Valor", "Unidad"]
            for c_idx, h in enumerate(headers, start=1):
                cell = ws.cell(row=5, column=c_idx, value=h)
                cell.fill = fill_header; cell.font = font_header
                cell.alignment = Alignment(horizontal="center")

            inp = self.data_export["inputs"]
            datos_in = [
                ("Tipo de Estructura de Entrada", "Proyecto", inp["Tipo Proyecto"], "-"),
                ("Caudal del Diseño",             "Q",        inp["Q"],             "m³/s"),
                ("Tirante Supercrítico",          "y₁",       inp["y1"],            "m"),
                ("Temperatura del Fluido",        "T",        inp["Temp"],          "°C"),
                ("Altitud del Proyecto",          "A",        inp["Alt"],           "msnm"),
                ("Ancho Solera Azud",             "Bᵣ",       inp["Br"],            "m"),
                ("Diámetro de Colector",          "D₀",       inp["D0"],            "m"),
            ]

            r_curr = 6
            for item in datos_in:
                ws.cell(row=r_curr, column=1, value=item[0]).font = font_regular
                ws.cell(row=r_curr, column=2, value=item[1]).font = font_bold
                v_cell = ws.cell(row=r_curr, column=3, value=item[2]); v_cell.font = font_regular
                if isinstance(item[2], float): v_cell.number_format = "0.00"
                ws.cell(row=r_curr, column=4, value=item[3]).font = font_regular
                for col in range(1, 5): ws.cell(row=r_curr, column=col).border = thin
                r_curr += 1

            r_curr += 1
            ws.cell(row=r_curr, column=1, value="2. RESULTADOS DEL DIMENSIONAMIENTO").font = font_section
            ws.merge_cells(start_row=r_curr, start_column=1, end_row=r_curr, end_column=4)
            ws.cell(row=r_curr, column=1).fill = fill_section

            r_curr += 1
            for c_idx, h in enumerate(headers, start=1):
                cell = ws.cell(row=r_curr, column=c_idx, value=h)
                cell.fill = fill_header; cell.font = font_header
                cell.alignment = Alignment(horizontal="center")

            r_curr += 1
            res = self.data_export["results"]
            datos_out = [
                ("Ancho del Cuenco",            "Wb",   res["Wb"],    "m"),
                ("Ancho Rápida Entrada",        "Wb₁",  res["Wb1"],   "m"),
                ("Ancho Sección Central",       "Wb₂",  res["Wb2"],   "m"),
                ("Ancho Salida",                "Wb₃",  res["Wb3"],   "m"),
                ("Velocidad Supercrítica",      "V₁",   res["V1"],    "m/s"),
                ("Número de Froude",            "Fr₁",  res["Fr1"],   "-"),
                ("Tirante Crítico",             "y꜀",   res["yc"],    "m"),
                ("Tirante Conjugado",           "y₂",   res["d2"],    "m"),
                ("Factor Corrector",            "C",    res["C"],     "-"),
                ("Transición de Entrada",       "Lₐ",   res["LA"],    "m"),
                ("Longitud Cuenco SAF",         "L_B",  res["LB"],    "m"),
                ("Ubicación Bloques",           "Dₑₛ",  res["Des"],   "m"),
                ("N.° Bloques Rápida",          "n꜀♭",  res["ncb"],   "u"),
                ("Ancho Bloques Rápida",        "w꜀♭",  res["wcb"],   "m"),
                ("Altura Bloques Rápida",       "h꜀♭",  res["hcb"],   "m"),
                ("N.° Bloques Fondo",           "nբ♭",  res["nfb"],   "u"),
                ("Ancho Bloques Fondo",         "wբ♭",  res["wfb"],   "m"),
                ("Altura Bloques Fondo",        "hբ♭",  res["hfb"],   "m"),
                ("Espesor Cresta",              "e",    res["e"],     "m"),
                ("Distancia a Pared",           "dᵣₑₐₗ",res["dreal"],"m"),
                ("Altura Umbral de Salida",     "h₄",   res["h4"],    "m"),
                ("Altura de Paredes",           "H₆",   res["h6"],    "m"),
                ("Borde Libre",                 "z_bl", res["z_bl"],  "m"),
                ("Índice de Cavitación",        "σ",    res["sigma"], "-"),
                ("Diagnóstico Cavitación", "Estado", self.data_export["cav_text"], "—"),
            ]

            for item in datos_out:
                ws.cell(row=r_curr, column=1, value=item[0]).font = font_regular
                ws.cell(row=r_curr, column=2, value=item[1]).font = font_bold
                v_cell = ws.cell(row=r_curr, column=3, value=item[2])
                if isinstance(item[2], float):
                    v_cell.font = font_regular
                    v_cell.number_format = "0.0000" if item[1] == "σ" else "0.00"
                elif isinstance(item[2], int):
                    v_cell.font = font_regular; v_cell.number_format = "0"
                else:
                    v_cell.font = font_bold
                ws.cell(row=r_curr, column=4, value=item[3]).font = font_regular
                if r_curr % 2 == 0:
                    for col in range(1, 5): ws.cell(row=r_curr, column=col).fill = fill_zebra
                for col in range(1, 5): ws.cell(row=r_curr, column=col).border = thin
                r_curr += 1

            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                ws.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 3, 11)
            ws.column_dimensions['A'].width = 38
            ws.column_dimensions['C'].width = 45

            wb.save(path)
            messagebox.showinfo("Éxito", "Memoria técnica exportada correctamente a Excel.")
        except Exception as e:
            messagebox.showerror("Error", f"No se pudo guardar la hoja de cálculo.\nDetalle: {e}")

# ------------------------------------------------------------------------------
# VISTA: MODULO COMPLETO - CUENCO TIPO USBR VI
# ------------------------------------------------------------------------------
class InterfaceUSBRVI(ctk.CTkFrame):
    def __init__(self, master, control_navegacion, **kwargs):
        super().__init__(master, **kwargs)
        self.configure(fg_color="#0F172A")
        self.control_navegacion = control_navegacion
        self.inputs  = {}
        self.results = {}
        self.checks  = {}
        self.status_labels = {}
        self.datos_para_exportar = None
        self._img_refs = []

        self.ruta_epn      = resolver_ruta(os.path.join("Imagenes", "Escuela_Politécnica_Nacional.png"))
        self.ruta_fica     = resolver_ruta(os.path.join("Imagenes", "logo_fica.png"))
        self.ruta_planta   = resolver_ruta(os.path.join("Imagenes", "Vista en planta_USBR_VI.png"))
        self.ruta_elevacion= resolver_ruta(os.path.join("Imagenes", "Sección A-A_USBR_VI.png"))

        self.init_module_ui()

    def init_module_ui(self):
        header = ctk.CTkFrame(self, fg_color="#1E293B", corner_radius=0)
        header.pack(fill="x", padx=0, pady=0)
        content_head = ctk.CTkFrame(header, fg_color="transparent")
        content_head.pack(pady=8, fill="x")

        self.btn_volver = ctk.CTkButton(
            content_head, text="⬅ Menú Principal", width=140, height=32,
            fg_color="#334155", hover_color="#475569", font=("Segoe UI", 11, "bold"),
            command=self.control_navegacion.mostrar_inicio)
        self.btn_volver.pack(side="left", padx=20)

        self.load_img(content_head, self.ruta_epn, (45, 45), "left")
        titles_frame = ctk.CTkFrame(content_head, fg_color="transparent")
        titles_frame.pack(side="left", padx=15, expand=True)
        ctk.CTkLabel(titles_frame, text="DISEÑO DE CUENCO DISIPADOR USBR TIPO VI",
                     font=("Segoe UI", 18, "bold"), text_color="#F8FAFC").pack()
        ctk.CTkLabel(titles_frame,
                     text="Dimensionamiento hidráulico de cámaras e infraestructura de disipación por impacto.",
                     font=("Segoe UI", 11, "italic"), text_color="#94A3B8").pack()
        self.load_img(content_head, self.ruta_fica, (110, 35), "right")

        body = ctk.CTkFrame(self, fg_color="transparent")
        body.pack(fill="both", expand=True, padx=15, pady=10)

        panels = ctk.CTkFrame(body, fg_color="transparent")
        panels.pack(fill="both", expand=True)

        # ── SECCIÓN IZQUIERDA (RESPIRA AUTOMÁTICAMENTE SIN WIDTH RESTRICTIVO) ──
        left_side = ctk.CTkFrame(panels, fg_color="transparent")
        left_side.pack(side="left", fill="both", expand=True, padx=(0, 15))
        self.left_scroll = ctk.CTkScrollableFrame(left_side, fg_color="#1E293B")
        self.left_scroll.pack(fill="both", expand=True)

        # ── SECCIÓN DERECHA (FIJA PARA LOS ESQUEMAS GRÁFICOS) ──
        right_side = ctk.CTkFrame(panels, fg_color="transparent", width=490)
        right_side.pack(side="right", fill="both", expand=False)
        right_side.pack_propagate(False)
        self.right_scroll = ctk.CTkScrollableFrame(right_side, fg_color="#1E293B")
        self.right_scroll.pack(fill="both", expand=True)

        ctk.CTkLabel(self.left_scroll, text="Datos de Entrada:",
                     font=("Segoe UI", 14, "bold"),
                     text_color="#38BDF8").pack(anchor="w", padx=15, pady=(8, 4))
        self.setup_inputs()

        ctk.CTkLabel(self.left_scroll, text="Resultados del Dimensionamiento:",
                     font=("Segoe UI", 14, "bold"),
                     text_color="#38BDF8").pack(anchor="w", padx=15, pady=(12, 4))
        self.setup_results_table()

        ctk.CTkLabel(self.right_scroll, text="Esquema de Detalle Geométrico USBR VI:",
                     font=("Segoe UI", 13, "bold"),
                     text_color="#38BDF8").pack(anchor="w", padx=10, pady=(5, 2))
        self.setup_graphics()

        # ── PIE DE PÁGINA ─────────────────────────────────────────────────────
        footer = ctk.CTkFrame(self, fg_color="#1E293B", height=80, corner_radius=0, border_width=1, border_color="#334155")
        footer.pack(fill="x", side="bottom")
        footer.pack_propagate(False)

        info_frame = ctk.CTkFrame(footer, fg_color="transparent")
        info_frame.pack(side="left", padx=25, pady=5)
        ctk.CTkLabel(info_frame, text="Facultad de Ingeniería Civil y Ambiental — EPN",
                     font=("Segoe UI", 11, "bold"), text_color="#F8FAFC").pack(anchor="w")
        ctk.CTkLabel(info_frame, text="Trabajo de Integración Curricular  |  Disipador USBR VI",
                     font=("Segoe UI", 10), text_color="#94A3B8").pack(anchor="w")

        btn_frame = ctk.CTkFrame(footer, fg_color="transparent")
        btn_frame.pack(side="right", padx=25, pady=20)
        ctk.CTkButton(btn_frame, text="Calcular", command=self.calcular,
                      fg_color="#0EA5E9", hover_color="#0284C7", text_color="#0F172A", width=100, height=32,
                      font=("Segoe UI", 12, "bold")).pack(side="left", padx=5)
        ctk.CTkButton(btn_frame, text="Limpiar", command=self.limpiar,
                      fg_color="#475569", hover_color="#334155", text_color="#F8FAFC", width=90, height=32).pack(side="left", padx=5)
        ctk.CTkButton(btn_frame, text="Exportar Excel", command=self.exportar,
                      fg_color="#16A34A", hover_color="#15803D", text_color="#FFFFFF", width=130, height=32,
                      font=("Segoe UI", 11, "bold")).pack(side="left", padx=5)

    def load_img(self, parent, path, size, side):
        if os.path.exists(path):
            try:
                img = ctk.CTkImage(light_image=Image.open(path), dark_image=Image.open(path), size=size)
                self._img_refs.append(img)
                ctk.CTkLabel(parent, image=img, text="").pack(side=side, padx=5)
            except Exception:
                pass

    def setup_inputs(self):
        frame = ctk.CTkFrame(self.left_scroll, fg_color="transparent")
        frame.pack(fill="x", pady=2)

        self.tipo_sec = tk.IntVar(value=1)
        sel_frame = ctk.CTkFrame(frame, fg_color="transparent")
        sel_frame.pack(fill="x", padx=10, pady=5)
        ctk.CTkRadioButton(sel_frame, text="Sección Rectangular", variable=self.tipo_sec, value=1,
                           command=self.toggle_entries,
                           font=("Segoe UI", 11, "bold"), text_color="#38BDF8").pack(side="left", padx=15)
        ctk.CTkRadioButton(sel_frame, text="Sección Circular", variable=self.tipo_sec, value=2,
                           command=self.toggle_entries,
                           font=("Segoe UI", 11, "bold"), text_color="#38BDF8").pack(side="left", padx=15)

        grid_in = ctk.CTkFrame(frame, fg_color="transparent")
        grid_in.pack(fill="x", padx=10, pady=(0, 6))
        for _ci in range(3):
            grid_in.grid_columnconfigure(_ci, weight=1)

        fields_info = [
            ("Caudal",      "Q", "m³/s", 0, 0),
            ("n Manning",   "n", "-",    0, 1),
            ("Pendiente",   "S", "m/m",  0, 2),
            ("Ancho base",  "b", "m",    1, 0),
            ("Diámetro",    "D", "m",    1, 1),
            ("Temperatura", "T", "°C",   1, 2),
            ("Altitud",     "A", "msnm", 2, 0),
        ]

        for name, sym, unit, row, col in fields_info:
            cell = ctk.CTkFrame(grid_in, fg_color="transparent")
            cell.grid(row=row, column=col, padx=6, pady=5, sticky="w")
            
            # Ajustado ancho de etiqueta a 115 y unidad a 45 para comodidad visual
            ctk.CTkLabel(cell, text=f"{name} ({sym}):", width=115,
                         anchor="w", font=("Segoe UI", 11), text_color="#E2E8F0").pack(side="left")
            ent = ctk.CTkEntry(cell, width=75, height=24, fg_color="#0F172A", text_color="#F8FAFC", border_color="#334155", justify="center")
            ent.pack(side="left", padx=2)
            ctk.CTkLabel(cell, text=unit, width=45, anchor="w",
                         font=("Segoe UI", 10), text_color="#64748B").pack(side="left", padx=2)
            self.inputs[sym] = ent

        self.toggle_entries()

    def toggle_entries(self):
        for k in ["Q", "n", "S", "T", "A"]: self.inputs[k].configure(state="normal")
        if self.tipo_sec.get() == 1:
            self.inputs["b"].configure(state="normal",   fg_color="#0F172A")
            self.inputs["D"].configure(state="disabled", fg_color="#1E293B")
        else:
            self.inputs["b"].configure(state="disabled", fg_color="#1E293B")
            self.inputs["D"].configure(state="normal",   fg_color="#0F172A")

    def setup_results_table(self):
        master_frame = ctk.CTkFrame(self.left_scroll, fg_color="transparent")
        master_frame.pack(fill="both", expand=True, pady=2, ipady=4)

        ctk.CTkLabel(master_frame, text="Verificaciones Límite de Diseño",
                     font=("Segoe UI", 12, "bold"), text_color="#F8FAFC").pack(anchor="w", padx=12, pady=(6, 2))
        check_container = ctk.CTkFrame(master_frame, fg_color="transparent")
        check_container.pack(fill="x", padx=12, pady=2)

        for k, text in [
            ("V", "Velocidad de aproximación admisible V ≤ 15.24 m/s:"),
            ("Q", "Caudal máximo normado de descarga Q ≤ 11.33 m³/s:")
        ]:
            f = ctk.CTkFrame(check_container, fg_color="transparent")
            f.pack(fill="x", pady=2)
            ctk.CTkLabel(f, text=text, font=("Segoe UI", 11), text_color="#CBD5E1").pack(side="left")
            self.checks[k] = ctk.CTkLabel(f, text="---",
                                           font=("Segoe UI", 11, "bold"), text_color="#64748B")
            self.checks[k].pack(side="left", padx=12)

        res_config = [
            ("1. Flujo de Aproximación", [
                ("Tirante normal", "Yn", "m"),
                ("Área mojada",    "A",  "m²"),
                ("Velocidad",      "V",  "m/s"),
                ("Número Froude",  "Fr", "-")
            ]),
            ("2. Geometría Cuenco USBR VI", [
                ("Ancho cuenco",    "W",  "m"),
                ("Longitud cuenco", "L",  "m"),
                ("Espesor frontal", "f",  "m"),
                ("Transición",      "e",  "m"),
                ("Altura total",    "H",  "m"),
                ("Lado bloque",     "a",  "m"),
                ("Ancho deflector", "b",  "m"),
                ("Contrafuerte",    "c",  "m"),
                ("Espesor pantalla","t",  "m"),
                ("Diámetro Roca",   "Dr", "m")
            ]),
            ("3. Cavitación — Ingreso de la rápida", [
                ("Índice cavitación", "sigma", "-")
            ])
        ]

        for sec_name, items in res_config:
            ctk.CTkLabel(master_frame, text=sec_name,
                         font=("Segoe UI", 12, "bold"), text_color="#7DD3FC",
                         anchor="w").pack(fill="x", padx=12, pady=(8, 2))
            container = ctk.CTkFrame(master_frame, fg_color="transparent")
            container.pack(fill="x", padx=12, pady=2)

            # SE SUBIÓ EL MINSIZE A 255 PARA EVITAR QUE COLISIONEN LAS COLUMNAS EN EL GRID
            for col_idx in range(3):
                container.grid_columnconfigure(col_idx, weight=1, minsize=255)

            for i, (name, sym, unit) in enumerate(items):
                r, c = divmod(i, 3)
                cell = ctk.CTkFrame(container, fg_color="transparent")
                cell.grid(row=r, column=c, padx=4, pady=4, sticky="w")

                # Ancho incrementado a 135 para que el nombre quepa y no empuje las unidades
                ctk.CTkLabel(cell, text=f"{name} [{sym}]:", width=135,
                             anchor="w", font=("Segoe UI", 10.5), text_color="#CBD5E1").pack(side="left")
                res_box = ctk.CTkLabel(cell, text="0.00", width=68, height=22,
                                        fg_color="#0F172A", text_color="#38BDF8", corner_radius=4,
                                        font=("Segoe UI", 11, "bold"))
                res_box.pack(side="left", padx=2)
                
                # Ancho holgado de 40 para la unidad métrica
                ctk.CTkLabel(cell, text=unit, width=40, anchor="w",
                             font=("Segoe UI", 10), text_color="#64748B").pack(side="left", padx=2)
                self.results[sym] = res_box

            if "Cavitación" in sec_name:
                cav_f = ctk.CTkFrame(master_frame, fg_color="transparent")
                cav_f.pack(fill="x", padx=12, pady=(4, 6))
                ctk.CTkLabel(cav_f, text="Diagnóstico de Cavitación:",
                             font=("Segoe UI", 11, "bold"), text_color="#F8FAFC").pack(side="left")
                self.status_labels["Cav"] = ctk.CTkLabel(
                    cav_f, text="Evaluar diseño numérico",
                    font=("Segoe UI", 11, "bold"), text_color="#64748B")
                self.status_labels["Cav"].pack(side="left", padx=12)

    def setup_graphics(self):
        frame = ctk.CTkFrame(self.right_scroll, fg_color="transparent")
        frame.pack(fill="both", expand=True, pady=2)

        for title, path in [
            ("VISTA EN PLANTA INTERIOR",               self.ruta_planta),
            ("VISTA EN ELEVACIÓN / SECCIÓN TRANSVERSAL",  self.ruta_elevacion)
        ]:
            ctk.CTkLabel(frame, text=title, font=("Segoe UI", 11, "bold"),
                         text_color="#F8FAFC").pack(pady=(8, 2))
            self.render_image(frame, path)
            ctk.CTkLabel(frame, text="Adaptado de Beichley (1971)",
                         font=("Segoe UI", 9, "italic"),
                         text_color="#64748B").pack(pady=(0, 6), anchor="e", padx=15)

    def render_image(self, parent, path):
        if os.path.exists(path):
            try:
                ctk_img = ctk.CTkImage(light_image=Image.open(path), dark_image=Image.open(path), size=(470, 215))
                self._img_refs.append(ctk_img)
                ctk.CTkLabel(parent, text="", image=ctk_img).pack(pady=4, padx=12)
            except Exception:
                pass
        else:
            ctk.CTkLabel(parent, text="[Esquema geométrico normado USBR VI]",
                         height=110, width=410, text_color="#64748B",
                         font=("Segoe UI", 10, "italic"),
                         fg_color="#1E293B", corner_radius=8).pack(pady=6)

    def get_geo_logic(self, y_val, tipo, b_reg, D_reg):
        if tipo == 1:
            A = b_reg * y_val; P = b_reg + 2.0 * y_val; T = b_reg; theta = None
        else:
            y_s   = max(1e-6, min(y_val, D_reg - 1e-6))
            theta = 2.0 * np.arccos(1.0 - (2.0 * y_s / D_reg))
            A     = (D_reg**2 / 8.0) * (theta - np.sin(theta))
            P     = (D_reg / 2.0) * theta
            T     = D_reg * np.sin(theta / 2.0)
        R = A / P if P > 0 else 0
        return A, P, R, T, theta

    def f_manning(self, y_val, Q_reg, n_reg, S_reg, tipo, b_reg, D_reg):
        A, _, R, _, _ = self.get_geo_logic(y_val, tipo, b_reg, D_reg)
        if A <= 0: return -Q_reg
        return (1.0 / n_reg) * A * (R**(2.0/3.0)) * np.sqrt(S_reg) - Q_reg

    def calcular(self):
        if self.tipo_sec.get() == 0:
            messagebox.showwarning("Atención", "Seleccione el tipo de sección de conducción.")
            return
        try:
            Q        = float(self.inputs["Q"].get().replace(',', '.'))
            n_m      = float(self.inputs["n"].get().replace(',', '.'))
            S_m      = float(self.inputs["S"].get().replace(',', '.'))
            temp_agua= float(self.inputs["T"].get().replace(',', '.'))
            altitud  = float(self.inputs["A"].get().replace(',', '.'))

            t_sec = self.tipo_sec.get()
            b_s   = float(self.inputs["b"].get().replace(',', '.')) if t_sec == 1 else 0
            D_d   = float(self.inputs["D"].get().replace(',', '.')) if t_sec == 2 else 0
            g     = 9.81

            y_calc = 0.2 * D_d if t_sec == 2 else 0.5; tol = 1e-9
            for _ in range(200):
                h_step = 1e-8
                f  = self.f_manning(y_calc, Q, n_m, S_m, t_sec, b_s, D_d)
                df = (self.f_manning(y_calc + h_step, Q, n_m, S_m, t_sec, b_s, D_d) - f) / h_step
                if abs(df) < 1e-12: df = 1e-12
                paso   = 0.5 * (f / df)
                y_new  = y_calc - paso
                if t_sec == 2: y_new = max(1e-4, min(y_new, 0.85 * D_d))
                else:          y_new = max(1e-3, y_new)
                if abs(y_new - y_calc) < tol: y_calc = y_new; break
                y_calc = y_new

            y_norm = round(y_calc, 2)
            A_f, P_f, R_f, T_f, _ = self.get_geo_logic(y_norm, t_sec, b_s, D_d)
            V_t = Q / A_f

            v_cumple = V_t <= 15.24
            q_cumple = Q   <= 11.33

            self.checks["V"].configure(
                text="SÍ CUMPLE" if v_cumple else "EXCEDE LÍMITE",
                text_color="#4ADE80" if v_cumple else "#EF4444")
            self.checks["Q"].configure(
                text="SÍ CUMPLE" if q_cumple else "EXCEDE LÍMITE",
                text_color="#4ADE80" if q_cumple else "#EF4444")

            if not v_cumple or not q_cumple:
                for r in self.results.values(): r.configure(text="0.00")
                self.status_labels["Cav"].configure(text="Diseño Fuera de Rango Normado", text_color="#EF4444")
                self.datos_para_exportar = None
                messagebox.showerror(
                    "Criterio de Diseño Excedido",
                    f"La estructura supera las envolventes hidráulicas USBR Tipo VI.\n\n"
                    f"Caudal: {Q:.2f} m³/s (Máx: 11.33 m³/s)\n"
                    f"Velocidad: {V_t:.2f} m/s (Máx: 15.24 m/s)")
                return

            Fr     = V_t / np.sqrt(g * (A_f / T_f))
            Q_cfs  = Q * 35.3147
            W_ft   = (1.47 * (Q_cfs**0.4) + 1.79 * (Q_cfs**0.4)) / 2.0
            m_conv = 0.3048

            H_atm    = calcular_presion_atmosferica(altitud)
            H_v      = calcular_presion_vapor(temp_agua)
            carga_vel= (V_t**2) / (2.0 * g)
            sigma_val= (H_atm - H_v) / carga_vel if carga_vel > 0 else 0

            if   sigma_val >= 1.0: estado_cav, color_cav = "ESTRUCTURA HIDRÁULICAMENTE SEGURA ✔", "#4ADE80"
            elif sigma_val >= 0.6: estado_cav, color_cav = "PROBABILIDAD DE CAVITACIÓN MODERADA ⚠", "#FBBF24"
            else:                  estado_cav, color_cav = "CRÍTICO — ALTO RIESGO DE CAVITACIÓN ✘", "#EF4444"

            res_map = {
                "Yn": round(y_norm,          2), "A":  round(A_f,               2),
                "V":  round(V_t,             2), "Fr": round(Fr,                2),
                "W":  round(W_ft * m_conv,   2), "L":  round(1.333 * W_ft * m_conv, 2),
                "f":  round(W_ft / 6.0 * m_conv, 2), "e": round(W_ft / 12.0 * m_conv, 2),
                "H":  round(0.75 * W_ft * m_conv, 2), "a": round(0.50 * W_ft * m_conv, 2),
                "b":  round(0.375 * W_ft * m_conv, 2), "c": round(0.50 * W_ft * m_conv, 2),
                "t":  round(W_ft / 12.0 * m_conv, 2), "Dr": round(W_ft / 20.0 * m_conv, 2),
                "sigma":  round(sigma_val, 4)
            }

            for k, v in res_map.items():
                if k in self.results:
                    self.results[k].configure(text=f"{v:.4f}" if k == "sigma" else f"{v:.2f}")

            self.status_labels["Cav"].configure(text=estado_cav, text_color=color_cav)

            self.datos_para_exportar = {
                "inputs": {
                    "Tipo Sección": "Rectangular" if t_sec == 1 else "Circular",
                    "Q": Q, "n": n_m, "S": S_m,
                    "b": b_s, "D": D_d, "T": temp_agua, "A": altitud},
                "results": res_map, "cav_diag": estado_cav
            }

        except Exception as e:
            messagebox.showerror("Error Funcional", f"Verifique las variables numéricas de entrada ingresadas.\nDetalle: {e}")

    def limpiar(self):
        for e in self.inputs.values():
            old = e.cget("state"); e.configure(state="normal"); e.delete(0, tk.END); e.configure(state=old)
        for r in self.results.values(): r.configure(text="0.00")
        for c in self.checks.values():  c.configure(text="---", text_color="#64748B")
        self.status_labels["Cav"].configure(text="Evaluar diseño numérico", text_color="#64748B")
        self.datos_para_exportar = None

    def exportar(self):
        if not self.datos_para_exportar:
            messagebox.showwarning("Atención", "Primero realice un cálculo válido.")
            return
        path = filedialog.asksaveasfilename(defaultextension=".xlsx",
                                            filetypes=[("Archivos Excel", "*.xlsx")])
        if not path: return
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Resultados USBR VI"
            ws.views.sheetView[0].showGridLines = True

            fill_header  = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            fill_section = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
            fill_zebra   = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

            font_title   = Font(name="Segoe UI", size=14, bold=True, color="1E293B")
            font_section = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            font_header  = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
            font_bold    = Font(name="Segoe UI", size=10, bold=True)
            font_regular = Font(name="Segoe UI", size=10)

            thin = Border(
                left=Side(style='thin', color='CBD5E1'), right=Side(style='thin', color='CBD5E1'),
                top=Side(style='thin', color='CBD5E1'), bottom=Side(style='thin', color='CBD5E1'))

            ws["A1"] = "MEMORIA TÉCNICA DE DISEÑO HIDRÁULICO - USBR TIPO VI"
            ws["A1"].font = font_title

            ws["A4"] = "1. DATOS INICIALES DE ENTRADA"
            ws.merge_cells("A4:D4"); ws["A4"].fill = fill_section; ws["A4"].font = font_section

            headers = ["Parámetro Técnico", "Símbolo", "Valor Numérico", "Unidad"]
            for c_idx, h in enumerate(headers, start=1):
                cell = ws.cell(row=5, column=c_idx, value=h)
                cell.fill = fill_header; cell.font = font_header
                cell.alignment = Alignment(horizontal="center")

            inp = self.datos_para_exportar["inputs"]
            datos_in = [
                ("Geometría de Conducción", "Sección", inp["Tipo Sección"], "-"),
                ("Caudal Máximo de Diseño",  "Q",       inp["Q"],           "m³/s"),
                ("Coeficiente de Manning",   "n",       inp["n"],           "-"),
                ("Pendiente de Canal",       "S",       inp["S"],           "m/m"),
                ("Ancho de Solera",          "b",       inp["b"],           "m"),
                ("Diámetro de Conducción",   "D",       inp["D"],           "m"),
                ("Temperatura del fluido",   "T",       inp["T"],           "°C"),
                ("Altitud del Proyecto",     "A",       inp["A"],           "msnm"),
            ]

            r_curr = 6
            for item in datos_in:
                ws.cell(row=r_curr, column=1, value=item[0]).font = font_regular
                ws.cell(row=r_curr, column=2, value=item[1]).font = font_bold
                v_cell = ws.cell(row=r_curr, column=3, value=item[2]); v_cell.font = font_regular
                if isinstance(item[2], float): v_cell.number_format = "0.00"
                ws.cell(row=r_curr, column=4, value=item[3]).font = font_regular
                for col in range(1, 5): ws.cell(row=r_curr, column=col).border = thin
                r_curr += 1

            r_curr += 1
            ws.cell(row=r_curr, column=1,
                    value="2. DIMENSIONAMIENTO GEOMÉTRICO Y SEGURIDAD").font = font_section
            ws.merge_cells(start_row=r_curr, start_column=1, end_row=r_curr, end_column=4)
            ws.cell(row=r_curr, column=1).fill = fill_section

            r_curr += 1
            for c_idx, h in enumerate(headers, start=1):
                cell = ws.cell(row=r_curr, column=c_idx, value=h)
                cell.fill = fill_header; cell.font = font_header
                cell.alignment = Alignment(horizontal="center")

            r_curr += 1
            res = self.datos_para_exportar["results"]
            nombres_largos = {
                "Yn": "Tirante Normal", "A": "Área Mojada", "V": "Velocidad de Entrada",
                "Fr": "Número de Froude", "W": "Ancho del Cuenco", "L": "Longitud del Cuenco",
                "f": "Espesor Frontal", "e": "Transición de Entrada", "H": "Altura Cuenco",
                "a": "Dimensión Lado a", "b": "Ancho Deflector", "c": "Contrafuerte",
                "t": "Espesor de Pantalla", "Dr": "Diámetro de Roca", "sigma": "Índice de Cavitación",
            }
            unidades = {
                "Yn":"m","A":"m2","V":"m/s","Fr":"-","W":"m","L":"m","f":"m","e":"m",
                "H":"m","a":"m","b":"m","c":"m","t":"m","Dr":"m","sigma":"-",
            }

            for k, v in res.items():
                ws.cell(row=r_curr, column=1, value=nombres_largos.get(k, k)).font = font_regular
                ws.cell(row=r_curr, column=2, value=k).font = font_bold
                v_cell = ws.cell(row=r_curr, column=3, value=v); v_cell.font = font_regular
                v_cell.number_format = "0.0000" if k == "sigma" else "0.00"
                ws.cell(row=r_curr, column=4, value=unidades.get(k, "")).font = font_regular
                if r_curr % 2 == 0:
                    for col in range(1, 5): ws.cell(row=r_curr, column=col).fill = fill_zebra
                for col in range(1, 5): ws.cell(row=r_curr, column=col).border = thin
                r_curr += 1

            ws.cell(row=r_curr, column=1, value="Diagnóstico Operativo").font = font_regular
            ws.cell(row=r_curr, column=2, value="Estado").font = font_bold
            ws.cell(row=r_curr, column=3, value=self.datos_para_exportar["cav_diag"]).font = font_bold
            ws.cell(row=r_curr, column=4, value="Diagnóstico").font = font_regular
            for col in range(1, 5): ws.cell(row=r_curr, column=col).border = thin

            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                ws.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 3, 12)
            ws.column_dimensions["A"].width = 38

            wb.save(path)
            messagebox.showinfo("Éxito", "Memoria técnica exportada correctamente a Excel.")
        except Exception as e:
            messagebox.showerror("Error de Escritura", f"No se pudo completar el guardado del archivo.\nDetalle: {e}")

# ------------------------------------------------------------------------------
# CONTROLADOR MAESTRO 
# ------------------------------------------------------------------------------
class AplicacionPrincipal(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title("Hidra Dis Pro - Cuencos Disipadores | EPN FICA")
        self.geometry("1450x920")
        self.minsize(1250, 780)
        self.grid_rowconfigure(0, weight=1)
        self.grid_columnconfigure(0, weight=1)

        self.pantalla_inicio  = PantallaInicio(master=self,  control_navegacion=self)
        self.pantalla_saf     = InterfaceSAF(master=self,    control_navegacion=self)
        self.pantalla_usbr_vi = InterfaceUSBRVI(master=self, control_navegacion=self)

        self.mostrar_inicio()

    def ocultar_vistas(self):
        self.pantalla_inicio.grid_forget()
        self.pantalla_saf.grid_forget()
        self.pantalla_usbr_vi.grid_forget()

    def mostrar_inicio(self):
        self.ocultar_vistas()
        self.pantalla_inicio.grid(row=0, column=0, sticky="nsew")

    def mostrar_saf(self):
        self.ocultar_vistas()
        self.pantalla_saf.grid(row=0, column=0, sticky="nsew")

    def mostrar_usbr_vi(self):
        self.ocultar_vistas()
        self.pantalla_usbr_vi.grid(row=0, column=0, sticky="nsew")


if __name__ == "__main__":
    app = AplicacionPrincipal()
    app.mainloop()