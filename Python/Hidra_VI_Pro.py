import customtkinter as ctk
from tkinter import messagebox, filedialog
from PIL import Image
import os
import sys
from datetime import datetime
import tkinter as tk
import numpy as np
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# ------------------------------------------------------------------------------
# PARTE 1: CONFIGURACIÓN GENERAL 
# ------------------------------------------------------------------------------
def resolver_ruta(ruta_relativa):
    """Resuelve rutas de recursos tanto en modo .py como en .exe (PyInstaller)."""
    try:
        base_path = sys._MEIPASS
    except AttributeError:
        base_path = os.path.abspath(".")
    return os.path.join(base_path, ruta_relativa)

# --- PALETA DE COLORES Y ESTILOS ---
ctk.set_appearance_mode("light")
AZUL_FONDO = "#D6E4F0" 
COLOR_BORDE = "#B1D0E0"
COLOR_ACCENTO = "#1A374D"
COLOR_RESULTADO = "#d32f2f"

# ------------------------------------------------------------------------------
# PARTE 2: MOTOR MATEMÁTICO E HIDRÁULICO 
# ------------------------------------------------------------------------------
def calcular_presion_vapor(temperatura_c: float) -> float:
    """Calcula la presión de vapor del agua (Hv) en m.c.a. utilizando la ecuación de Antoine."""
    p_mmhg = 10**(8.07131 - (1730.63 / (temperatura_c + 233.426)))
    return 0.0136 * p_mmhg

def calcular_presion_atmosferica(altitud_msnm: float) -> float:
    """Calcula la presión atmosférica (H_atm) en m.c.a. según la altitud local."""
    return 10.33 - (altitud_msnm / 900.0)

# ------------------------------------------------------------------------------
# PARTE 3: INTERFAZ GRÁFICA DE USUARIO (GUI)
# ------------------------------------------------------------------------------
class USBR_VI_Designer_Final(ctk.CTk):
    def __init__(self):
        super().__init__()

        self.title("USBR VI Designer Pro - Alex Ante")
        self.geometry("1450x900") 
        self.configure(fg_color=AZUL_FONDO)

        # RUTAS DE RECURSOS 
        self.ruta_epn       = resolver_ruta(os.path.join("Imagenes", "Escuela_Politécnica_Nacional.png"))
        self.ruta_fica      = resolver_ruta(os.path.join("Imagenes", "logo_fica.png"))
        self.ruta_planta    = resolver_ruta(os.path.join("Imagenes", "Vista en planta_USBR_VI.png"))
        self.ruta_elevacion = resolver_ruta(os.path.join("Imagenes", "Sección A-A_USBR_VI.png"))

        self._img_refs = []
        self.inputs = {}
        self.results = {}
        self.checks = {}
        self.status_labels = {}
        self.datos_para_exportar = None
        self.init_ui()

    def load_img(self, parent, path, size, side):
        try:
            if os.path.exists(path):
                ctk_img = ctk.CTkImage(light_image=Image.open(path), size=size)
                self._img_refs.append(ctk_img)
                ctk.CTkLabel(parent, image=ctk_img, text="").pack(side=side, padx=5)
        except Exception:
            pass

    def init_ui(self):
        # ENCABEZADO 
        header_main = ctk.CTkFrame(self, fg_color=AZUL_FONDO, corner_radius=0)
        header_main.pack(fill="x", padx=20, pady=(5, 2))
        
        header_logos = ctk.CTkFrame(header_main, fg_color="transparent")
        header_logos.pack(fill="x")
        self.load_img(header_logos, self.ruta_epn, (60, 60), "left")
        
        title_container = ctk.CTkFrame(header_logos, fg_color="transparent")
        title_container.pack(side="left", expand=True)
        ctk.CTkLabel(title_container, text="DISEÑO DE CUENCO DISIPADOR USBR TIPO VI - Hidra_VI Pro", 
                     font=("Segoe UI", 20, "bold"), text_color=COLOR_ACCENTO).pack(pady=0)
        ctk.CTkLabel(title_container, 
                     text="Este programa permite el cálculo y dimensionamiento de cuencos disipadores de impacto USBR Tipo VI", 
                     font=("Segoe UI", 12, "italic"), text_color="#406882").pack(pady=0)
        
        self.load_img(header_logos, self.ruta_fica, (130, 40), "right")

        # CUERPO PRINCIPAL 
        main_scroll = ctk.CTkScrollableFrame(self, fg_color=AZUL_FONDO, scrollbar_button_color="#406882")
        main_scroll.pack(fill="both", expand=True, padx=15, pady=2)

        content_frame = ctk.CTkFrame(main_scroll, fg_color="transparent")
        content_frame.pack(fill="both", expand=True)

        left_panel = ctk.CTkFrame(content_frame, fg_color="transparent", width=740)
        left_panel.pack(side="left", fill="both", expand=True, padx=(0, 5))
        
        right_panel = ctk.CTkFrame(content_frame, fg_color="transparent", width=640)
        right_panel.pack(side="right", fill="both", expand=True)

        ctk.CTkLabel(left_panel, text=" Datos: ", font=("Segoe UI", 13, "bold"), text_color=COLOR_ACCENTO).pack(anchor="w", padx=2, pady=(2, 0))
        self.setup_inputs(left_panel)

        ctk.CTkLabel(left_panel, text=" Resultados: ", font=("Segoe UI", 13, "bold"), text_color=COLOR_ACCENTO).pack(anchor="w", padx=2, pady=(4, 0))
        self.setup_results_table(left_panel)

        ctk.CTkLabel(right_panel, text=" Esquema : ", font=("Segoe UI", 13, "bold"), text_color=COLOR_ACCENTO).pack(anchor="w", padx=5, pady=(2, 0))
        self.setup_graphics(right_panel)

        # BARRA DE PIE DE PÁGINA (BOTONES DE ACCIÓN) 
        footer = ctk.CTkFrame(self, fg_color=AZUL_FONDO, height=75, corner_radius=0, border_width=1, border_color=COLOR_BORDE)
        footer.pack(fill="x", side="bottom")
        footer.pack_propagate(False)

        info_frame = ctk.CTkFrame(footer, fg_color="transparent")
        info_frame.pack(side="left", padx=25, pady=4)
        
        ctk.CTkLabel(info_frame, text="Facultad de Ingeniería Civil y Ambiental - EPN", font=("Segoe UI", 10, "bold"), text_color=COLOR_ACCENTO).pack(anchor="w", pady=0)
        ctk.CTkLabel(info_frame, text="Trabajo de Integración Curricular", font=("Segoe UI", 10, "bold"), text_color="#406882").pack(anchor="w", pady=0)
        
        fecha_hoy = datetime.now().strftime("%d/%m/%Y")
        ctk.CTkLabel(info_frame, text=f"Autor: Alex Ante  |  Fecha: {fecha_hoy}", font=("Segoe UI", 10, "bold"), text_color=COLOR_ACCENTO).pack(anchor="w", pady=(1, 0))

        btn_frame = ctk.CTkFrame(footer, fg_color="transparent")
        btn_frame.pack(side="right", padx=25, pady=15)
        ctk.CTkButton(btn_frame, text="Calcular", command=self.calcular, fg_color=COLOR_ACCENTO, width=95, height=30, font=("Segoe UI", 11, "bold")).pack(side="left", padx=4)
        ctk.CTkButton(btn_frame, text="Limpiar", command=self.limpiar, fg_color="#698396", width=85, height=30).pack(side="left", padx=4)
        ctk.CTkButton(btn_frame, text="Exportar Excel", command=self.exportar, fg_color="#2E7D32", width=125, height=30, font=("Segoe UI", 11, "bold")).pack(side="left", padx=4)

    def setup_inputs(self, parent):
        frame = ctk.CTkFrame(parent, fg_color=AZUL_FONDO, corner_radius=8, border_width=1, border_color=COLOR_BORDE)
        frame.pack(fill="x", pady=2) 
        
        self.tipo_sec = tk.IntVar(value=1)
        sel_frame = ctk.CTkFrame(frame, fg_color="transparent")
        sel_frame.pack(fill="x", padx=10, pady=4)
        ctk.CTkRadioButton(sel_frame, text="Rectangular", variable=self.tipo_sec, value=1, command=self.toggle_entries, font=("Segoe UI", 11, "bold"), text_color=COLOR_ACCENTO).pack(side="left", padx=10)
        ctk.CTkRadioButton(sel_frame, text="Circular", variable=self.tipo_sec, value=2, command=self.toggle_entries, font=("Segoe UI", 11, "bold"), text_color=COLOR_ACCENTO).pack(side="left", padx=10)
        
        grid_in = ctk.CTkFrame(frame, fg_color="transparent")
        grid_in.pack(fill="x", padx=10, pady=(0, 6))
        
        #  Inyección de Altitud (A) 
        fields_info = [
            ("Caudal", "Q", "m³/s", 0, 0), 
            ("n Manning", "n", "-", 0, 1), 
            ("Pendiente", "S", "m/m", 0, 2), 
            ("Ancho base", "b", "m", 1, 0), 
            ("Diámetro", "D", "m", 1, 1),
            ("Temperatura", "T", "°C", 1, 2),
            ("Altitud", "A", "msnm", 2, 0)
        ]
        
        for name, sym, unit, row, col in fields_info:
            cell = ctk.CTkFrame(grid_in, fg_color="transparent")
            cell.grid(row=row, column=col, padx=5, pady=2, sticky="w")
            ctk.CTkLabel(cell, text=f"{name} ({sym}):", width=90, anchor="w", font=("Segoe UI", 10.5)).pack(side="left")
            ent = ctk.CTkEntry(cell, width=65, height=22, fg_color="white", justify="center")
            ent.pack(side="left", padx=2)
            self.inputs[sym] = ent
            ctk.CTkLabel(cell, text=unit, width=35, anchor="w", font=("Segoe UI", 9.5), text_color="gray").pack(side="left")
            
        self.toggle_entries()

    def toggle_entries(self):
        for k in ["Q", "n", "S", "T", "A"]: self.inputs[k].configure(state="normal")
        if self.tipo_sec.get() == 1:
            self.inputs["b"].configure(state="normal", fg_color="white")
            self.inputs["D"].configure(state="disabled", fg_color="#E0E0E0")
        else:
            self.inputs["b"].configure(state="disabled", fg_color="#E0E0E0")
            self.inputs["D"].configure(state="normal", fg_color="white")

    def setup_results_table(self, parent):
        master_frame = ctk.CTkFrame(parent, fg_color=AZUL_FONDO, corner_radius=8, border_width=1, border_color=COLOR_BORDE)
        master_frame.pack(fill="both", expand=True, pady=2, ipady=4)
        
        ctk.CTkLabel(master_frame, text="Verificaciones Normativas", font=("Segoe UI", 11, "bold"), text_color=COLOR_ACCENTO).pack(anchor="w", padx=12, pady=(4, 1))
        check_container = ctk.CTkFrame(master_frame, fg_color="transparent")
        check_container.pack(fill="x", padx=12, pady=1)
        
        for k, text in [("V", "Velocidad de aproximación V ≤ 15.24 m/s:"), ("Q", "Caudal de diseño de descarga Q ≤ 11.33 m³/s:")]:
            f = ctk.CTkFrame(check_container, fg_color="transparent")
            f.pack(fill="x", pady=1)
            ctk.CTkLabel(f, text=text, font=("Segoe UI", 10.5)).pack(side="left")
            self.checks[k] = ctk.CTkLabel(f, text="---", font=("Segoe UI", 10.5, "bold"), text_color="gray")
            self.checks[k].pack(side="left", padx=10)

        res_config = [
            ("1. Flujo de Aproximación", [("Tirante normal", "Yn", "m"), ("Área mojada", "A", "m²"), ("Velocidad", "V", "m/s"), ("Número Froude", "Fr", "-")]),
            ("2. Geometría Cuenco USBR VI", [
                ("Ancho cuenco", "W", "m"), ("Longitud cuenco", "L", "m"), ("Espesor frontal", "f", "m"), 
                ("Transición", "e", "m"), ("Altura total", "H", "m"), ("Lado bloque", "a", "m"), 
                ("Ancho deflector", "b", "m"), ("Contrafuerte", "c", "m"), ("Espesor pantalla", "t", "m"), 
                ("Diámetro Roca", "Dr", "m")
            ]),
            ("3. Cavitación - Ingreso de la rápida", [("Índice cavitación", "σ", "-")])
        ]
        
        for sec_name, items in res_config:
            ctk.CTkLabel(master_frame, text=sec_name, font=("Segoe UI", 11, "bold"), text_color="#406882", anchor="w").pack(fill="x", padx=12, pady=(6, 1))
            container = ctk.CTkFrame(master_frame, fg_color="transparent")
            container.pack(fill="x", padx=12, pady=1)
            
            for i, (name, sym, unit) in enumerate(items):
                r, c = divmod(i, 3) 
                row = ctk.CTkFrame(container, fg_color="transparent")
                row.grid(row=r, column=c, padx=3, pady=2, sticky="w")
                
                ctk.CTkLabel(row, text=f"{name} [{sym}]:", width=115, anchor="w", font=("Segoe UI", 10)).pack(side="left")
                res_box = ctk.CTkLabel(row, text="0.00", width=65, height=20, fg_color="white", corner_radius=4, font=("Segoe UI", 10.5, "bold"), text_color=COLOR_RESULTADO)
                res_box.pack(side="left", padx=2)
                self.results[sym] = res_box
                ctk.CTkLabel(row, text=unit, width=25, anchor="w", font=("Segoe UI", 9.5), text_color="gray").pack(side="left")

            if "Cavitación" in sec_name:
                status_cav = ctk.CTkFrame(master_frame, fg_color="transparent")
                status_cav.pack(fill="x", padx=12, pady=(2, 4))
                ctk.CTkLabel(status_cav, text="Diagnóstico de Cavitación:", font=("Segoe UI", 10, "bold"), text_color=COLOR_ACCENTO).pack(side="left")
                self.status_labels["Cav"] = ctk.CTkLabel(status_cav, text="Evaluar diseño numérico", font=("Segoe UI", 10, "bold"), text_color="gray")
                self.status_labels["Cav"].pack(side="left", padx=10)

    def setup_graphics(self, parent):
        frame = ctk.CTkFrame(parent, fg_color=AZUL_FONDO, corner_radius=8, border_width=1, border_color=COLOR_BORDE)
        frame.pack(fill="both", expand=True, pady=2)
        for title, path in [("VISTA EN PLANTA ", self.ruta_planta), ("VISTA EN ELEVACIÓN / SECCIÓN CORTE", self.ruta_elevacion)]:
            ctk.CTkLabel(frame, text=title, font=("Segoe UI", 11, "bold"), text_color=COLOR_ACCENTO).pack(pady=(4, 1))
            self.render_image(frame, path)
            ctk.CTkLabel(frame, text="Adaptado Beichely (1971)", 
                         font=("Segoe UI", 8, "italic"), text_color="#555555").pack(pady=(0, 4), anchor="e", padx=15)

    def render_image(self, parent, path):
        try:
            if os.path.exists(path):
                ctk_img = ctk.CTkImage(light_image=Image.open(path), size=(480, 220))
                self._img_refs.append(ctk_img)
                ctk.CTkLabel(parent, text="", image=ctk_img).pack(pady=2, padx=10)
            else:
                ctk.CTkLabel(parent, text="[Esquema de diseño USBR VI — archivo no encontrado en recursos]",
                             height=100, width=420, text_color="gray",
                             font=("Segoe UI", 9, "italic"), fg_color="#E0E0E0",
                             corner_radius=6).pack(pady=5)
        except Exception:
            pass

    def get_geo_logic(self, y_val, tipo, b_reg, D_reg):
        if tipo == 1:
            A = b_reg * y_val; P = b_reg + 2 * y_val; T = b_reg; theta = None
        else:
            y_s = max(1e-6, min(y_val, D_reg - 1e-6))
            theta = 2 * np.arccos(1 - (2 * y_s / D_reg))
            A = (D_reg**2 / 8) * (theta - np.sin(theta)); P = (D_reg / 2) * theta; T = D_reg * np.sin(theta / 2)
        R = A / P if P > 0 else 0
        return A, P, R, T, theta

    def f_manning(self, y_val, Q_reg, n_reg, S_reg, tipo, b_reg, D_reg):
        A, _, R, _, _ = self.get_geo_logic(y_val, tipo, b_reg, D_reg)
        if A <= 0: return -Q_reg
        return (1/n_reg) * A * (R**(2/3)) * np.sqrt(S_reg) - Q_reg

# ------------------------------------------------------------------------------
# PARTE 4: PROCESAMIENTO LÓGICO Y CÁLCULO 
# ------------------------------------------------------------------------------
    def calcular(self):
        if self.tipo_sec.get() == 0:
            messagebox.showwarning("Atención", "Seleccione el tipo de sección.")
            return
        try:
            Q = float(self.inputs["Q"].get().replace(',', '.'))
            n_m = float(self.inputs["n"].get().replace(',', '.'))
            S_m = float(self.inputs["S"].get().replace(',', '.'))
            temp_agua = float(self.inputs["T"].get().replace(',', '.')) 
            altitud = float(self.inputs["A"].get().replace(',', '.')) # CORRECCIÓN: Lectura de Altitud
            
            t_sec = self.tipo_sec.get()
            b_s = float(self.inputs["b"].get().replace(',', '.')) if t_sec == 1 else 0
            D_d = float(self.inputs["D"].get().replace(',', '.')) if t_sec == 2 else 0
            g = 9.81
            
            # Algoritmo de Newton-Raphson
            y_calc = 0.1 * D_d if t_sec == 2 else 0.5; tol = 1e-9
            for _ in range(200):
                h_step = 1e-8; f = self.f_manning(y_calc, Q, n_m, S_m, t_sec, b_s, D_d)
                df = (self.f_manning(y_calc + h_step, Q, n_m, S_m, t_sec, b_s, D_d) - f) / h_step
                paso = 0.6 * (f / df); y_new = y_calc - paso
                if t_sec == 2: y_new = max(1e-4, min(y_new, 0.85 * D_d))
                else: y_new = max(1e-3, y_new)
                if abs(y_new - y_calc) < tol: y_calc = y_new; break
                y_calc = y_new
                
            y_norm = round(y_calc, 2); A_f, P_f, R_f, T_f, _ = self.get_geo_logic(y_norm, t_sec, b_s, D_d)
            V_t = Q / A_f
            
            # VERIFICACIÓN DE LÍMITES ESTRICTOS DE LA NORMA USBR TIPO VI 
            v_cumple = V_t <= 15.24
            q_cumple = Q <= 11.33 
            
            # Actualizar visualmente los indicadores normativos superiores
            self.checks["V"].configure(text="SÍ CUMPLE" if v_cumple else "NO CUMPLE", text_color="green" if v_cumple else "red")
            self.checks["Q"].configure(text="SÍ CUMPLE" if q_cumple else "NO CUMPLE", text_color="green" if q_cumple else "red")
            
            # BLOQUEO DE SEGURIDAD CRÍTICO 
            if not v_cumple or not q_cumple:
                for r in self.results.values(): r.configure(text="0.00")
                self.status_labels["Cav"].configure(text="Diseño Fuera de Rango", text_color="gray")
                self.datos_para_exportar = None
                
                messagebox.showerror(
                    "Error de Diseño: Límites Superados", 
                    f"El diseño NO cumple con los criterios empíricos del cuenco de impacto USBR Tipo VI.\
\
"
                    f"Detalle de Falla:\
"
                    f" - Caudal calculado: {Q:.2f} m³/s (Máximo permitido: 11.33 m³/s)\
"
                    f" - Velocidad calculada: {V_t:.2f} m/s (Máximo permitido: 15.24 m/s)\
\
"
                    f"El proceso se ha detenido y no se pueden generar las dimensiones de la estructura."
                )
                return
                
            Fr = V_t / np.sqrt(g * (A_f / T_f))
            Q_cfs = Q * 35.3147
            W_ft = (1.47 * (Q_cfs**0.4) + 1.79 * (Q_cfs**0.4)) / 2
            m_conv = 0.3048
            
            # Cálculo dinámico de la presión atmosférica según altitud
            H_atm = calcular_presion_atmosferica(altitud)
            H_v = calcular_presion_vapor(temp_agua) 
            carga_vel = (V_t**2) / (2 * g)
            sigma_val = (H_atm - H_v) / carga_vel if carga_vel > 0 else 0
            
            if sigma_val >= 1.0:
                estado_cav = "Condición Segura (No hay Cavitación)"
                color_cav = "green"
            elif 0.6 <= sigma_val < 1.0:
                estado_cav = "Riesgo Moderado (Hay Probabilidad de Cavitación)"
                color_cav = "orange"
            else:
                estado_cav = "¡Alto Riesgo! (Sí hay Cavitación Crítica)"
                color_cav = "red"

            res_map = {
                "Yn": round(y_norm, 2), "A": round(A_f, 2), "V": round(V_t, 2), "Fr": round(Fr, 2),
                "W": round(W_ft * m_conv, 2), "L": round((1.333 * W_ft) * m_conv, 2),
                "f": round((W_ft / 6) * m_conv, 2), "e": round((W_ft / 12) * m_conv, 2),
                "H": round((0.75 * W_ft) * m_conv, 2), "a": round((0.50 * W_ft) * m_conv, 2),
                "b": round((0.375 * W_ft) * m_conv, 2), "c": round((0.50 * W_ft) * m_conv, 2),
                "t": round((W_ft / 12) * m_conv, 2), "Dr": round((W_ft / 20) * m_conv, 2),
                "σ": round(sigma_val, 2)
            }
            
            for k, v in res_map.items(): 
                self.results[k].configure(text=f"{v:.4f}" if k == "σ" else f"{v:.2f}")
                
            self.status_labels["Cav"].configure(text=estado_cav, text_color=color_cav)
            
            self.datos_para_exportar = {
                "inputs": {"Tipo Sección": "Rectangular" if t_sec == 1 else "Circular", "Q": Q, "n": n_m, "S": S_m, "b": b_s, "D": D_d, "T": temp_agua, "A": altitud},
                "results": res_map,
                "cav_diag": estado_cav
            }
        except Exception as e: 
            messagebox.showerror("Error de Datos", f"Verifique que todos los campos contengan valores numéricos: {e}")

    def limpiar(self):
        for e in self.inputs.values():
            old_state = e.cget("state")
            e.configure(state="normal")
            e.delete(0, tk.END)
            e.configure(state=old_state)
        for r in self.results.values(): r.configure(text="0.00")
        for c in self.checks.values(): c.configure(text="---", text_color="gray")
        self.status_labels["Cav"].configure(text="Evaluar diseño numérico", text_color="gray")
        self.datos_para_exportar = None

# ------------------------------------------------------------------------------
# PARTE 5: EXPORTACIÓN AUTOMATIZADA A EXCEL
# ------------------------------------------------------------------------------
    def exportar(self):
        if not self.datos_para_exportar:
            messagebox.showwarning("Atención", "Primero debe realizar un cálculo válido normado.")
            return
        path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Archivos Excel", "*.xlsx")], title="Guardar Memoria Técnica")
        if not path: return
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Resultados USBR VI"
            ws.views.sheetView[0].showGridLines = True

            fill_header = PatternFill(start_color="1A374D", end_color="1A374D", fill_type="solid")
            fill_section = PatternFill(start_color="406882", end_color="406882", fill_type="solid")
            fill_zebra = PatternFill(start_color="F9FBFC", end_color="F9FBFC", fill_type="solid")
            
            font_title = Font(name="Segoe UI", size=14, bold=True, color="1A374D")
            font_section = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
            font_bold = Font(name="Segoe UI", size=10, bold=True)
            font_regular = Font(name="Segoe UI", size=10)

            thin_border = Border(
                left=Side(style='thin', color='B1D0E0'), right=Side(style='thin', color='B1D0E0'),
                top=Side(style='thin', color='B1D0E0'), bottom=Side(style='thin', color='B1D0E0')
            )

            ws["A1"] = "MEMORIA TÉCNICA DE DISEÑO HIDRÁULICO - USBR TIPO VI"
            ws["A1"].font = font_title
        
            # TABLA 1: DATOS INICIALES DE ENTRADA
            ws["A4"] = "1. DATOS INICIALES DE ENTRADA"
            ws.merge_cells("A4:D4"); ws["A4"].fill = fill_section; ws["A4"].font = font_section
            
            headers = ["Parámetro Técnico", "Símbolo", "Valor Numérico", "Unidad"]
            for c_idx, h in enumerate(headers, start=1):
                cell = ws.cell(row=5, column=c_idx, value=h)
                cell.fill = fill_header; cell.font = font_header; cell.alignment = Alignment(horizontal="center")

            inp = self.datos_para_exportar["inputs"]
            datos_in = [
                ("Geometría de Conducción", "Sección", inp["Tipo Sección"], "-"),
                ("Caudal Máximo de Diseño", "Q", inp["Q"], "m³/s"),
                ("Coeficiente de Manning", "n", inp["n"], "-"),
                ("Pendiente de Canal", "S", inp["S"], "m/m"),
                ("Ancho de Solera", "b", inp["b"], "m"),
                ("Diámetro de Conducción", "D", inp["D"], "m"),
                ("Temperatura del fluido", "T", inp["T"], "°C"),
                ("Altitud del Proyecto", "A", inp["A"], "msnm") # CORRECCIÓN: Inyección en Excel
            ]

            r_curr = 6
            for item in datos_in:
                ws.cell(row=r_curr, column=1, value=item[0]).font = font_regular
                ws.cell(row=r_curr, column=2, value=item[1]).font = font_bold
                v_cell = ws.cell(row=r_curr, column=3, value=item[2]); v_cell.font = font_regular
                if isinstance(item[2], float): v_cell.number_format = "0.00"
                ws.cell(row=r_curr, column=4, value=item[3]).font = font_regular
                for col in range(1, 5): ws.cell(row=r_curr, column=col).border = thin_border
                r_curr += 1

            # TABLA 2: RESULTADOS Y DIMENSIONAMIENTO
            r_curr += 1
            ws.cell(row=r_curr, column=1, value="2. DIMENSIONAMIENTO GEOMÉTRICO Y SEGURIDAD").font = font_section
            ws.merge_cells(start_row=r_curr, start_column=1, end_row=r_curr, end_column=4); ws.cell(row=r_curr, column=1).fill = fill_section
            
            r_curr += 1
            for c_idx, h in enumerate(headers, start=1):
                cell = ws.cell(row=r_curr, column=c_idx, value=h)
                cell.fill = fill_header; cell.font = font_header; cell.alignment = Alignment(horizontal="center")
            
            r_curr += 1
            res = self.datos_para_exportar["results"]
            nombres_largos = {
                "Yn": "Tirante Normal", "A": "Área Mojada", "V": "Velocidad de Entrada", "Fr": "Número de Froude", 
                "W": "Ancho del Cuenco", "L": "Longitud del Cuenco", "f": "Espesor Frontal Bloque", 
                "e": "Transición de Entrada", "H": "Altura Cuenco", "a": "Dimensión Lado 'a'", 
                "b": "Ancho Deflector de Fondo", "c": "Contrafuerte Posterior", "t": "Espesor de Pantalla", 
                "Dr": "Diámetro de Roca", "σ": "Índice de Cavitación"
            }
            unidades = {
                "Yn": "m", "A": "m²", "V": "m/s", "Fr": "-", "W": "m", "L": "m", "f": "m", 
                "e": "m", "H": "m", "a": "m", "b": "m", "c": "m", "t": "m", "Dr": "m", "σ": "-"
            }

            for k, v in res.items():
                ws.cell(row=r_curr, column=1, value=nombres_largos.get(k, k)).font = font_regular
                ws.cell(row=r_curr, column=2, value=k).font = font_bold
                v_cell = ws.cell(row=r_curr, column=3, value=v); v_cell.font = font_regular
                v_cell.number_format = "0.0000" if k == "σ" else "0.00"
                ws.cell(row=r_curr, column=4, value=unidades.get(k, "")).font = font_regular
                
                if r_curr % 2 == 0:
                    for col in range(1, 5): ws.cell(row=r_curr, column=col).fill = fill_zebra
                for col in range(1, 5): ws.cell(row=r_curr, column=col).border = thin_border
                r_curr += 1

            # Diagnóstico en reporte
            ws.cell(row=r_curr, column=1, value="Diagnóstico Operativo").font = font_regular
            ws.cell(row=r_curr, column=2, value="Estado").font = font_bold
            v_cell = ws.cell(row=r_curr, column=3, value=self.datos_para_exportar["cav_diag"]); v_cell.font = font_bold
            ws.cell(row=r_curr, column=4, value="Diagnóstico").font = font_regular
            for col in range(1, 5): ws.cell(row=r_curr, column=col).border = thin_border

            # Auto-ajuste de columnas
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
            ws.column_dimensions['A'].width = 38

            wb.save(path)
            messagebox.showinfo("Éxito", "Memoria técnica organizada por tablas y exportada correctamente a Excel.")
        except Exception as e:
            messagebox.showerror("Error de Archivo", f"No se pudo guardar la memoria.\
Detalle: {e}")

# ------------------------------------------------------------------------------
# EJECUCIÓN DE APLICACIÓN
# ------------------------------------------------------------------------------
if __name__ == "__main__":
    app = USBR_VI_Designer_Final()
    app.mainloop()